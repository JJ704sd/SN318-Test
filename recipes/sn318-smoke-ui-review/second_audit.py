from pathlib import Path
import re, json
import openpyxl

base=Path(__file__).parent
path=Path('./artifacts/smoke-ui/SN318_提测前冒烟与UI规范_用例复核.xlsx')
w=openpyxl.load_workbook(path,data_only=False)
rows={str(w['本批提测'].cell(r,1).value):r for r in list(range(40,62))+list(range(115,118))}
cases={w[s].cell(r,1).value:(s,r,'AF' if s=='UI检查' else 'V') for s,n in [('冒烟用例',93),('UI检查',45)] for r in range(6,n+1)}
missing=[]
for did,r in rows.items():
    text=w['本批提测'].cell(r,3).value
    for m in re.finditer(r'(SM|UI)(\d+)(?:[—–-](?:SM|UI)?(\d+))?',str(text)):
        for n in range(int(m[2]),int(m[3] or m[2])+1):
            cid=f'{m[1]}{n:02d}'
            s,cr,col=cases[cid]
            f=w[s][f'{col}{cr}'].value
            if not re.search(rf"'本批提测'!\$H\${r}(?!\d)",f):
                missing.append([did,cid,f])
print('Missing decision dependencies',json.dumps(missing,ensure_ascii=False))
print('Case SM55',w['冒烟用例']['V60'].value)
new={s:[[str(w[s].cell(r,c).value or '') for c in [1,2,3,10,11,12]] for r in range(6,w[s].max_row+1)] for s in ['冒烟用例','UI检查']}
(base/'second-cases.json').write_text(json.dumps(new,ensure_ascii=False,indent=2),encoding='utf-8')
current=[json.loads(l) for l in (base/'ui-second.ndjson').read_text(encoding='utf-8-sig').splitlines()]
first=[json.loads(l) for l in (base/'ui-live.ndjson').read_text(encoding='utf-8-sig').splitlines()]
assert current==first
assert len(cases)==128
assert len({c for c in cases})==128
live_ids={r['record_id'] for r in current if r['当前原型问题'] or r['Review 建议']}
file_ids={w['UI检查'].cell(r,28).value for r in range(6,46) if w['UI检查'].cell(r,28).value}
assert live_ids==file_ids
print('20 source Review IDs match, source unchanged, 128 unique cases')
print('SM52',w['冒烟用例']['J57'].value,w['冒烟用例']['K57'].value)
