import json, pathlib, hashlib
import openpyxl

root = pathlib.Path('.')
work = root / '.work/ui-audit-advice'
for name, path in [
    ('public', work / 'common-baseline-live.json'),
    ('product', work / 'ui-spec-live.json'),
    ('ui_old', root / '.work/sn318-research/ui-spec-v1.2.json'),
    ('ui_recent', root / '.work/sn318-smoke-ui-review/spec-live.json'),
]:
    data = json.loads(path.read_text(encoding='utf-8-sig'))
    doc = data.get('data', {}).get('document', {})
    print(name, 'ok', data.get('ok'), 'keys', list(doc))
    print(json.dumps({k:v for k,v in doc.items() if k not in ['content','comments','comment_aliases','blocks','styles','block_aliases']}, ensure_ascii=False)[:1500])
    content = doc.get('content', '')
    (work / (name + '-readable.txt')).write_text(content, encoding='utf-8')
    if name == 'public':
        print(content)

src = root / 'outputs/smoke-ui/SN318_提测前冒烟与UI规范_颜色层级对应_复查修正.xlsx'
wb = openpyxl.load_workbook(src, data_only=False)
print('sheets', [(s.title, s.max_row, s.max_column) for s in wb])
sheet = wb['UI检查']
headers = [c.value for c in sheet[5]]
print('headers', json.dumps(list(enumerate(headers, 1)), ensure_ascii=False))
cases = [dict(zip(headers, [c.value for c in row])) for row in list(sheet)[5:] if row[0].value]
(work / 'source-ui-cases.json').write_text(json.dumps(cases,ensure_ascii=False,indent=2),encoding='utf-8')
for c in cases:
    print(json.dumps({k:v for k,v in c.items() if k in headers[:3] or k in headers[18:20] or k in headers[25:29]}, ensure_ascii=False))
(work / 'source-hash.json').write_text(json.dumps({'path': str(src), 'sha256': hashlib.sha256(src.read_bytes()).hexdigest()}, ensure_ascii=False),encoding='utf-8')
