import copy
import hashlib
import html
import json
import posixpath
import zipfile
import sys
from collections import Counter
from pathlib import Path
import xml.etree.ElementTree as ET
import openpyxl

ROOTWORK=Path('./recipes/template-ui-reaudit-20261014')
THIRD='third' in sys.argv
SECOND='second' in sys.argv or THIRD
CONFIG=json.loads((ROOTWORK/('third-config.json' if THIRD else 'second-config.json')).read_text(encoding='utf-8')) if SECOND else None
WORK=Path(CONFIG['work']) if SECOND else ROOTWORK
OUT=Path('./artifacts/template-reaudit')
SRC=Path('./artifacts/weekly-ui/SN318_W2_UI与交互测试用例_按工单用例风格重写.xlsx')
ART=WORK/'artifact-export.xlsx'
DEST=OUT/'SN318_W2_UI与交互测试用例_40条实测复核修订_20261014.xlsx'
if SECOND:
    OUT,SRC,DEST=Path(CONFIG['out']),Path(CONFIG['source']),Path(CONFIG['destination'])
NS='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
RNS='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
ET.register_namespace('',NS)
ET.register_namespace('r',RNS)
Q=lambda tag: f'{{{NS}}}{tag}'
edits=json.loads((WORK/'edits.json').read_text(encoding='utf-8'))
cases=json.loads((WORK/'cases.json').read_text(encoding='utf-8'))
expected={(e['sheet'],e['cell']):e['value'] for e in edits}

def sheet_paths(z):
    wb=ET.fromstring(z.read('xl/workbook.xml'))
    rels=ET.fromstring(z.read('xl/_rels/workbook.xml.rels'))
    targets={e.attrib['Id']:e.attrib['Target'] for e in rels}
    return {e.attrib['name']:posixpath.normpath('xl/'+targets[e.attrib[f'{{{RNS}}}id']]) if not targets[e.attrib[f'{{{RNS}}}id']].startswith('/') else targets[e.attrib[f'{{{RNS}}}id']].lstrip('/') for e in wb.find(Q('sheets'))}

def value_node(z,cell):
    # Consume the value authored and exported by Artifact Tool. Re-encode only
    # edited string cells so the original shared-string/style tables and all
    # unrelated package parts can remain byte-for-byte intact.
    if cell.attrib.get('t')=='inlineStr':
        return copy.deepcopy(cell.find(Q('is')))
    if cell.attrib.get('t')=='s':
        strings=ET.fromstring(z.read('xl/sharedStrings.xml'))
        node=ET.Element(Q('is'))
        for child in strings[int(cell.find(Q('v')).text)]:
            node.append(copy.deepcopy(child))
        return node
    if cell.attrib.get('t')=='str':
        node=ET.Element(Q('is'))
        text=ET.SubElement(node,Q('t'))
        text.attrib['{http://www.w3.org/XML/1998/namespace}space']='preserve'
        text.text=cell.find(Q('v')).text
        return node
    raise AssertionError(f'Expected authored text value at {cell.attrib}')

with zipfile.ZipFile(SRC) as original, zipfile.ZipFile(ART) as artifact:
    opaths,apaths=sheet_paths(original),sheet_paths(artifact)
    replacements={}
    for name in {e['sheet'] for e in edits}:
        base=ET.fromstring(original.read(opaths[name]))
        authored=ET.fromstring(artifact.read(apaths[name]))
        acells={e.attrib['r']:e for e in authored.iter(Q('c'))}
        bcells={e.attrib['r']:e for e in base.iter(Q('c'))}
        for edit in [e for e in edits if e['sheet']==name]:
            dest=bcells.get(edit['cell'])
            if dest is None:
                assert THIRD and name==CONFIG['manualColumn']['sheet'] and edit['cell']=='J1'
                row=base.find(Q('sheetData')).find(f'{Q("row")}[@r="1"]')
                dest=ET.SubElement(row,Q('c'),{'r':'J1','s':bcells['I1'].attrib.get('s','0')})
                bcells['J1']=dest
            for child in list(dest):
                if child.tag in {Q('v'),Q('is'),Q('f')}:
                    dest.remove(child)
            dest.attrib['t']='inlineStr'
            dest.append(value_node(artifact,acells[edit['cell']]))
        if THIRD and name==CONFIG['manualColumn']['sheet']:
            # J values/format and expanded table were authored by Artifact Tool.
            # Translate the copied native I style to the original style IDs,
            # retaining every other workbook package part and existing style.
            data=base.find(Q('sheetData'))
            for n in range(2,42):
                authored_cell=acells.get(f'J{n}')
                if authored_cell is not None:
                    assert authored_cell.find(Q('v')) is None and authored_cell.find(Q('is')) is None
                row=data.find(f'{Q("row")}[@r="{n}"]')
                ET.SubElement(row,Q('c'),{'r':f'J{n}','s':bcells[f'I{n}'].attrib.get('s','0')})
            for row in data:
                if 'spans' in row.attrib:row.attrib['spans']='1:10'
            base.find(Q('dimension')).attrib['ref']=CONFIG['manualColumn']['range']
            acol=next(e for e in authored.find(Q('cols')) if int(e.attrib['min'])<=10<=int(e.attrib['max']))
            newcol=copy.deepcopy(acol);newcol.attrib.update({'min':'10','max':'10'})
            newcol.attrib.pop('style',None)
            base.find(Q('cols')).append(newcol)
        replacements[opaths[name]]=ET.tostring(base,encoding='utf-8',xml_declaration=True)
    if THIRD:
        def find_table(z,name):
            for p in z.namelist():
                if p.startswith('xl/tables/') and p.endswith('.xml'):
                    element=ET.fromstring(z.read(p))
                    if element.attrib.get('name')==name:return p,element
            raise AssertionError('Requested table missing')
        tpath,base_table=find_table(original,CONFIG['manualColumn']['tableName'])
        _,auth_table=find_table(artifact,CONFIG['manualColumn']['tableName'])
        assert auth_table.attrib['ref']==CONFIG['manualColumn']['range']
        authored_columns=auth_table.find(Q('tableColumns'))
        assert len(authored_columns)==10 and authored_columns[-1].attrib['name']==CONFIG['manualColumn']['header']
        base_table.attrib['ref']=auth_table.attrib['ref']
        base_table.find(Q('autoFilter')).attrib['ref']=auth_table.attrib['ref']
        cols=base_table.find(Q('tableColumns'));cols.attrib['count']='10';cols.append(copy.deepcopy(authored_columns[-1]))
        assert len({c.attrib['id'] for c in cols})==10
        replacements[tpath]=ET.tostring(base_table,encoding='utf-8',xml_declaration=True)
    OUT.mkdir(parents=True,exist_ok=True)
    with zipfile.ZipFile(DEST,'w',compression=zipfile.ZIP_DEFLATED,compresslevel=6) as target:
        for info in original.infolist():
            target.writestr(info,replacements.get(info.filename,original.read(info.filename)))

before,after=openpyxl.load_workbook(SRC),openpyxl.load_workbook(DEST)
assert before.sheetnames==after.sheetnames
actual_changes=[]
for bs,as_ in zip(before,after):
    is_manual=THIRD and bs.title==CONFIG['manualColumn']['sheet']
    assert (bs.max_row,bs.max_column+(1 if is_manual else 0))==(as_.max_row,as_.max_column)
    assert list(bs.merged_cells.ranges)==list(as_.merged_cells.ranges)
    assert bs.freeze_panes==as_.freeze_panes
    assert bs.data_validations.to_tree().attrib==as_.data_validations.to_tree().attrib
    assert ET.tostring(bs.data_validations.to_tree())==ET.tostring(as_.data_validations.to_tree())
    assert list(bs.tables)==list(as_.tables)
    for table in bs.tables:
        if is_manual and table==CONFIG['manualColumn']['tableName']:
            modified=copy.deepcopy(as_.tables[table])
            modified.ref=bs.tables[table].ref
            modified.autoFilter.ref=bs.tables[table].autoFilter.ref
            modified.tableColumns=modified.tableColumns[:-1]
            assert ET.tostring(bs.tables[table].to_tree())==ET.tostring(modified.to_tree())
        else:assert ET.tostring(bs.tables[table].to_tree())==ET.tostring(as_.tables[table].to_tree())
    for name,dim in bs.row_dimensions.items():
        assert dict(dim)==dict(as_.row_dimensions[name]) and dim.height==as_.row_dimensions[name].height
    for name,dim in bs.column_dimensions.items():
        assert dict(dim)==dict(as_.column_dimensions[name]) and dim.width==as_.column_dimensions[name].width
    for row in bs:
        for cell in row:
            new=as_[cell.coordinate]
            key=(bs.title,cell.coordinate)
            assert cell._style==new._style, key
            assert cell.hyperlink==new.hyperlink, key
            assert str(cell.comment)==str(new.comment), key
            if key in expected:
                assert new.value==expected[key], (key,new.value,expected[key])
            else:
                assert cell.value==new.value,(key,cell.value,new.value)
            if cell.value!=new.value:
                actual_changes.append({'sheet':bs.title,'cell':cell.coordinate,'before':cell.value,'after':new.value})
    if is_manual:
        assert as_.column_dimensions['J'].width==CONFIG['manualColumn']['width']
        for n in range(1,42):
            added=as_[f'J{n}'];prior=bs[f'I{n}']
            assert added._style==prior._style
            assert added.value==(CONFIG['manualColumn']['header'] if n==1 else None)
        actual_changes.append({'sheet':bs.title,'cell':'J1','before':None,'after':as_['J1'].value})
assert len(actual_changes)==len(edits)
with zipfile.ZipFile(SRC) as a,zipfile.ZipFile(DEST) as b:
    assert a.namelist()==b.namelist()
    changed_parts=[p for p in a.namelist() if a.read(p)!=b.read(p)]
    assert set(changed_parts)==set(replacements)

evidence=OUT/'evidence'
file_index={}
raw_index={}
for f in sorted(evidence.iterdir()):
    prefix=f.stem.split('-')[0]
    if f.suffix=='.jpg':
        file_index.setdefault(prefix,[]).append(f.name)
    elif f.suffix in ['.json','.txt']:
        raw_index.setdefault(prefix,[]).append(f.name)
for case in cases:
    for prefix in case['evidence']:
        assert prefix in file_index or prefix in raw_index, (case['id'],prefix)
assert len(cases)==40
counts=dict(Counter(c['status'] for c in cases))
assert counts==(CONFIG['expectedCounts'] if SECOND else {'不通过':11,'通过':12,'受阻':17})
assert sum(c['status']!=c.get('previousStatus',c['oldStatus']) for c in cases)==1
if SECOND:
    assert len({c['id'] for c in cases})==40 and all(c['secondAudit'] and c['expected'] and c['steps'] and c['prerequisites'] for c in cases)
    assert all('/demo/form' not in c['result'] for c in cases)

def esc(s): return html.escape(str(s))
blocks=[]
for c in cases:
    imgs=''.join(f'<a href="evidence/{esc(file)}"><img loading="lazy" src="evidence/{esc(file)}" alt="{esc(prefix)}"><span>{esc(file)} · {"本轮16条核实" if prefix.startswith("B") else "第二轮" if prefix.startswith("A") else "首轮10/14"}</span></a>' for prefix in c['evidence'] for file in file_index.get(prefix,[]))
    links=''.join(f'<a href="evidence/{esc(file)}">{esc(file)}</a><br>' for prefix in c['evidence'] for file in raw_index.get(prefix,[]))
    context=f'<p><b>{esc(c["secondMode"])}：</b>{esc(c["secondAudit"])}</p><details><summary>原前置条件、步骤与预期</summary><pre>前置条件\n{esc(c["prerequisites"])}\n操作步骤\n{esc(c["steps"])}\n预期结果\n{esc(c["expected"])}</pre></details>' if SECOND else ''
    history=''
    if SECOND:
        historical=OUT/'historical-20261013'
        hfiles=[f for f in sorted(historical.iterdir()) if f.stem.split('-')[0] in c['historicalEvidence']]
        himgs=''.join(f'<a href="historical-20261013/{esc(f.name)}"><img loading="lazy" src="historical-20261013/{esc(f.name)}" alt="历史10/13 {esc(f.name)}"><span>历史10/13：{esc(f.name)}</span></a>' for f in hfiles if f.suffix=='.jpg')
        history=f'<details><summary>历史10/13参考（本轮未重新执行）</summary><div class="images">{himgs}</div>'+''.join(f'<a href="historical-20261013/{esc(f.name)}">历史 {esc(f.name)}</a><br>' for f in hfiles if f.suffix=='.json')+'</details>' if hfiles else ''
    previous=f'<h3>上一版实际结果</h3><pre>{esc(c["previousResult"])}</pre>' if SECOND else ''
    scope_attr=f' data-scope="{"current" if c.get("inCurrentScope") else "previous"}"' if THIRD else ''
    blocks.append(f'<article id="{c["id"]}" data-status="{esc(c["status"])}"{scope_attr}><h2>{c["id"]} · {esc(c["name"])} <small>{esc(c.get("previousStatus",c["oldStatus"]))} → {esc(c["status"])}</small></h2>{context}<p>{esc(c["reason"])}</p><pre>{esc(c["result"])}</pre><div class="images">{imgs}</div>{history}<details><summary>DOM、量测与历史实际结果</summary>{links}{previous}<h3>原文件实际结果</h3><pre>{esc(c["originalResult"])}</pre></details></article>')
report='''<!doctype html><html lang="zh-CN"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>SN318 通用模板40条复核证据 · 2026-10-14</title><style>body{margin:0;background:#f4f6f9;color:#1a2130;font:15px/1.65 system-ui,"Microsoft YaHei",sans-serif}main{max-width:1150px;margin:auto;padding:24px}h1{font-size:24px}h2{font-size:18px;margin:0}small{display:block;font-weight:400;color:#666}article{background:white;padding:24px;border:1px solid #ddd;border-radius:12px;margin:20px 0}pre{white-space:pre-wrap;font:14px/1.7 inherit;overflow-wrap:anywhere}.images{display:grid;grid-template-columns:repeat(auto-fit,minmax(280px,1fr));gap:12px}.images img{width:100%;border:1px solid #ddd}.images span{display:block;font-size:12px}a{color:#046efc}button{padding:8px 14px;margin:4px;border:1px solid #ccc;background:white;border-radius:8px;cursor:pointer}header{position:sticky;top:0;background:#f4f6f9;padding:8px 0;z-index:1}summary{cursor:pointer}</style><main><h1>SN318 通用模板40条实测结果复核</h1><p>2026-10-14 · 本地 dev / Mock · 通过12、不通过11、受阻17。028由受阻改为不通过；014次错位置已修正，030末页已补测。原工作簿和业务76条保留。</p><p>规范正文 V1.2 / revision 5256 与10/13相同；UI Review指定视图及3.4、4.2、4.5.2规则由当前浏览器核对。R001–R010等历史参考未全部重新渲染。每条结果区分本次实测、历史证据及缺样例/策略的子步骤。</p><p>来源：<a href="https://example.invalid/source/516654b415;view=vewolLIwPP">UI Review指定视图</a> · <a href="https://example.invalid/source/05cc32fecc">通用设计规范V1.2</a> · <a href="sources/norm-live.json">本次规范正文快照</a></p><header><button onclick="filter('')">全部40条</button><button onclick="filter('不通过')">不通过11条</button><button onclick="filter('受阻')">受阻17条</button><button onclick="filter('通过')">通过12条</button></header>'''+''.join(blocks)+'''</main><script>function filter(v){document.querySelectorAll('article').forEach(e=>e.hidden=v&&e.dataset.status!==v)}</script></html>'''
if SECOND:
    report=report.replace('SN318 通用模板40条复核证据 · 2026-10-14','SN318 通用模板40条二次复核证据 · 2026-10-14')
    report=report.replace('SN318 通用模板40条实测结果复核','SN318 通用模板40条二次复核')
    report=report.replace('通过12、不通过11、受阻17。028由受阻改为不通过；014次错位置已修正，030末页已补测。原工作簿和业务76条保留。','通过12、不通过12、受阻16。019由上一版受阻改不通过；004链接颜色子项改不通过；表单真实路径纠正为/demo/list/form，032/033交互措辞与014修改清单次错坐标同步纠正。原文件、上一版和业务76条保留。')
    report=report.replace('规范正文 V1.2 / revision 5256 与10/13相同；UI Review指定视图及3.4、4.2、4.5.2规则由当前浏览器核对。R001–R010等历史参考未全部重新渲染。每条结果区分本次实测、历史证据及缺样例/策略的子步骤。','40条均逐项核对原前置、操作、预期与证据；重点分支补充实测，未宣称40条全部重新完整执行。第二轮规范正文仍revision 5256；3.1现行Brand Blue表明确链接#046EFC。Review指定视图再次核对24条及10/12 17:23更新时间，具体文本及3.4/4.2/4.5.2沿用首轮10/14截图。其他嵌入表及R001–R010仍为10/13历史，已分别打包。A为第二轮，V/S为首轮10/14，E/R为历史10/13。A06/V78瞬态量测与V79截图不是同时采样。')
    report=report.replace('不通过11条','不通过12条').replace('受阻17条','受阻16条')
    report=report.replace('<a href="sources/norm-live.json">本次规范正文快照</a>','<a href="sources/norm-second.json">第二轮规范正文快照</a> · <a href="sources/norm-live.json">首轮规范XML快照</a>')
if THIRD:
    report=report.replace('SN318 通用模板40条二次复核','SN318 原受阻16条核实')
    old_intro='通过12、不通过12、受阻16。019由上一版受阻改不通过；004链接颜色子项改不通过；表单真实路径纠正为/demo/list/form，032/033交互措辞与014修改清单次错坐标同步纠正。原文件、上一版和业务76条保留。'
    report=report.replace(old_intro,'本轮核实原受阻16条，011在固定300ms Mock范围完成并改通过；其余15条写明已测步骤及具体缺少条件。通用40条汇总：通过13、不通过12、受阻15。其他24条沿用上一版，未重复执行；业务76条与原文件保留。Excel新增J列“人工二次复核”，40行留空，由人工填结论、复核人、日期及意见。')
    report=report.replace('40条均逐项核对原前置、操作、预期与证据；重点分支补充实测，未宣称40条全部重新完整执行。第二轮规范正文仍revision 5256；','规范及Review来源沿用上一版10/14已读取的快照，本轮未重新逐记录读取。第二轮规范正文revision 5256；')
    report=report.replace('A为第二轮，V/S为首轮10/14，E/R为历史10/13。','B为本轮16条核实，A为二次复核，V/S为首轮10/14，E/R为历史10/13。B02/B05/B28/B40/B44为瞬态DOM；B48仅源码检查，不代替页面响应实测。')
    report=report.replace('B48仅源码检查，不代替页面响应实测。','B48仅源码检查，不代替页面响应实测。截图、DOM与量测依次采样，动态进度值以JSON为准，不保证逐帧同步。')
    report=report.replace('>通过12条</button>','>通过13条</button>').replace('>受阻16条</button>','>受阻15条</button>')
    report=report.replace('<header><button onclick="filter(\'\')">','<header><button onclick="filterScope()">本轮16条</button><button onclick="filter(\'\')">')
    report=report.replace('</script>','function filterScope(){document.querySelectorAll("article").forEach(e=>e.hidden=e.dataset.scope!=="current")}filterScope();</script>')
(OUT/'复核证据.html').write_text(report,encoding='utf-8')
(OUT/'sources').mkdir(exist_ok=True)
(OUT/'sources'/'norm-live.json').write_bytes((ROOTWORK/'norm-live.json').read_bytes())
sources={'norm':'V1.2, revision 5256, current XML equals 2026-10-13 snapshot','review':'Browser read specified view, 24 displayed records including blank row; no external writes','liveEmbeddedTables':['3.4','4.2','4.5.2'],'historicalReferences':'Other embedded tables and R001-R010 retained from 2026-10-13; not all rendered again','environment':'127.0.0.1:5173, dev/Mock; demo routes follow original case objects'}
if SECOND:
    sources.update({'secondNorm':'V1.2, revision 5256, read again with docs +fetch','secondLiveEmbeddedTable':'3.1 Brand Blue, links #046EFC, screenshot A30','secondReview':'A34: specified view 24 displayed records incl blank; 10/12 17:23; RV02 and batchbar text use first-pass S04/S07, not all record_ids refreshed','secondCoverage':'All 40 cases checked against original prerequisites/steps/expected and evidence; selected branches rerun; not 40 complete re-executions','historicalFolder':'historical-20261013; older evidence is labelled separately'})
if THIRD:sources.update({'currentCoverage':'All 16 previously blocked cases inspected; current executable branches tested; other 24 rows retained without rerun','currentEvidence':'B01-B48; transient DOM and source-only evidence are labelled','captureOrdering':'DOM, metrics, screenshot captured sequentially; dynamic progress frames are not synchronised','currentSources':'Feishu sources retained from second audit on 10/14; not re-read record by record this run','manualReviewColumn':'J, 人工二次复核, blank for 40 rows; no human verdict fabricated'})
(OUT/'sources'/'source-boundary.json').write_text(json.dumps(sources,ensure_ascii=False,indent=2),encoding='utf-8')
status_changes=[f'{c["id"]}:{c.get("previousStatus",c["oldStatus"])}→{c["status"]}' for c in cases if c['status']!=c.get('previousStatus',c['oldStatus'])]
validation={'passed':True,'cases':40,'counts':counts,'editedCells':len(actual_changes),'statusChanges':status_changes,'changedPackageParts':changed_parts,'unchangedBusinessCases':76,'preservedNativeFeatures':{'sheets':4,'tables':sum(len(s.tables) for s in before),'validations':sum(len(s.data_validations.dataValidation) for s in before),'mergedRanges':sum(len(s.merged_cells.ranges) for s in before),'styles':'all cells unchanged','rowColumnDimensions':'unchanged','unrelatedPackageParts':'byte-for-byte unchanged'},'sourceSha256':hashlib.sha256(SRC.read_bytes()).hexdigest(),'outputSha256':hashlib.sha256(DEST.read_bytes()).hexdigest(),'artifactExport':'authored with Artifact Tool; final package preserves original native features','evidencePhotos':sum(len(x) for x in file_index.values())}
if SECOND: validation.update({'secondAuditRecords':40,'editedResultCells':sum(e['sheet']=='通用模板UI与交互用例' and e['cell'].startswith('I') for e in edits),'secondPhotoCount':sum(len(v) for k,v in file_index.items() if k.startswith('A')),'secondCoverage':'evidence audit of 40, selected live re-tests, not 40 full executions'})
if THIRD:
    validation.update({'currentAuditRecords':16,'remainingBlocked':15,'completedCurrentCase':'UI-TPL-011','currentPhotoCount':sum(len(v) for k,v in file_index.items() if k.startswith('B')),'manualColumn':{'header':'人工二次复核','range':'J2:J41','blankRows':40,'tableRange':after[CONFIG['manualColumn']['sheet']].tables[CONFIG['manualColumn']['tableName']].ref},'currentScopeIds':CONFIG['auditScopeIds']})
    validation['preservedNativeFeatures']['styles']='all original cells unchanged; new J copies native I styles'
    validation['preservedNativeFeatures']['rowColumnDimensions']='original dimensions unchanged; added J width45'
    validation['secondAuditRecords']=40
(WORK/'validation.json').write_text(json.dumps(validation,ensure_ascii=False,indent=2),encoding='utf-8')
(WORK/'change-log.json').write_text(json.dumps(actual_changes,ensure_ascii=False,indent=2),encoding='utf-8')
(OUT/'复核索引.json').write_text(json.dumps({'validation':validation,'cases':cases},ensure_ascii=False,indent=2),encoding='utf-8')
archive_name='SN318_受阻16条核实_证据包_20261014.zip' if THIRD else 'SN318_通用模板40条_二次复核证据包_20261014.zip' if SECOND else 'SN318_通用模板40条_复核证据包_20261014.zip'
with zipfile.ZipFile(OUT/archive_name,'w',zipfile.ZIP_DEFLATED,compresslevel=6) as z:
    for name in ['复核证据.html','复核索引.json']:
        z.write(OUT/name,name)
    if THIRD:z.write(OUT/'本轮16条核实索引.json','本轮16条核实索引.json')
    for folder in ['evidence','sources']+(['historical-20261013'] if SECOND else []):
        for f in sorted((OUT/folder).rglob('*')):
            if f.is_file(): z.write(f,f.relative_to(OUT).as_posix())
print(json.dumps(validation,ensure_ascii=False,indent=2))
