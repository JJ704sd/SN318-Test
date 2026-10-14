import json,zipfile,hashlib
from pathlib import Path
import openpyxl
W=Path(__file__).parent
d=json.loads((W/'cases-data.json').read_text(encoding='utf-8'))
p=Path(d['out'])/'PCS_W2前端测试用例_按任务节奏细化.xlsx'
w=openpyxl.load_workbook(p);cache=openpyxl.load_workbook(p,data_only=True)
assert w.sheetnames==['本周执行','前端用例','工作包对应']
s=w['前端用例']
ids=[s.cell(r,1).value for r in range(6,103)]
assert len(set(ids))==97
assert sum(i.startswith('RV') for i in ids)==12
assert sum(i.startswith('FE') for i in ids)==47
assert sum(i.startswith('BP') for i in ids)==38
for r in range(6,103):
    assert s.cell(r,11).value=='待评审'
    assert s.cell(r,12).value==('不适用' if r<18 else '未执行')
    assert s.cell(r,13).value==('不适用' if r<18 else '未执行')
    assert s.cell(r,14).value in [None,''] and s.cell(r,15).value in [None,'']
    assert all(s.cell(r,c).value for c in [1,2,3,4,5,6,7,8,9,10])
    assert cache['前端用例'].cell(r,17).value=='待评审/依据'
packs=w['工作包对应']
assert {packs.cell(r,1).value for r in range(6,44)}=={r[0] for r in d['packages']}
for r in range(6,44):
    assert packs.cell(r,12).value==f'BP{packs.cell(r,1).value:02d}'
    for c in [10,11,12]:
        assert all(i in ids for i in packs.cell(r,c).value.split('、'))
assert [cache['本周执行'][c].value for c in ['B4','D4','F4','H4']]==[97,38,97,0]
assert s.tables['FrontendCases'].ref=='A5:Q102'
assert packs.tables['PackageMapping'].ref=='A5:N43'
assert s.freeze_panes=='C6' and packs.freeze_panes=='C6'
assert len(s.data_validations.dataValidation)>=2
errors=[(s.title,c.coordinate,c.value) for s in cache for row in s for c in row if c.data_type=='e']
assert not errors,errors
with zipfile.ZipFile(p) as z:assert z.testzip() is None
gate=json.loads((W/'gate-verification.json').read_text(encoding='utf-8'));assert all(gate.values())
result={'file':str(p),'bytes':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'cases':97,'designChecks':12,'componentCases':47,'packageCases':38,'mappedPackages':38,'formulaErrors':errors,'allResultsUnexecuted':True,'reviewAndDevEvidenceIndependentlyRequired':True,'savedFileReadback':'passed'}
(W/'verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(result,ensure_ascii=False))
