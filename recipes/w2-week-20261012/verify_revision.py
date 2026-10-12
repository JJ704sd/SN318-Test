import json,zipfile,hashlib
from pathlib import Path
from collections import Counter
import openpyxl

work=Path(__file__).parent
data=json.loads((work/'revision-data.json').read_text(encoding='utf-8'))
path=Path(data['out'])/'SN318_W2测试用例修订与本周任务.xlsx'
old=openpyxl.load_workbook(data['source'])
new=openpyxl.load_workbook(path)
cache=openpyxl.load_workbook(path,data_only=True)
changes=json.loads((work/'changes.json').read_text(encoding='utf-8'))
allowed={(x['sheet'],x['address']) for x in changes}
allowed.add(('UI走查','A57'))
diff=[];checked=0;formula_same=0;native=[]
for s in old:
    ns=new[s.title]
    for row in s:
        for c in row:
            if c.value is None: continue
            if (s.title,c.coordinate) in allowed: continue
            checked+=1
            if c.data_type=='f': formula_same+=1
            if c.value!=ns[c.coordinate].value: diff.append([s.title,c.coordinate,str(c.value)[:180],str(ns[c.coordinate].value)[:180]])
    for name in s.tables:
        t=s.tables[name]
        nt=ns.tables[name]
        expected={'UiCases':'A5:AM56','UIWalkthrough':'A6:Y57'}.get(name,t.ref)
        if nt.ref!=expected: native.append([s.title,name,t.ref,nt.ref])
    assert ns.freeze_panes==s.freeze_panes,(s.title,s.freeze_panes,ns.freeze_panes)
assert not diff, diff[:15]
assert not native, native
assert len(new['冒烟用例']['A6:A96'])==91
assert [new['UI检查'].cell(r,1).value for r in range(6,57)]==[f'UI{i:02d}' for i in range(1,52)]
assert [new['UI走查'].cell(r,1).value for r in range(7,58)]==[f'UI{i:02d}' for i in range(1,52)]
for r in [55,56]:
    assert new['UI检查'].cell(r,7).value=='未执行'
    assert new['UI检查'].cell(r,4).value=='待选择'
    assert f"'UI走查'!Y{r+1}" in new['UI检查'].cell(r,9).value
for r in range(6,82):
    assert new['W2用例'].cell(r,12).value=='未执行'
assert len({new['W2用例'].cell(r,1).value for r in range(6,82)})==76
assert set(new['W2范围'].cell(r,1).value for r in range(6,44))==set(p['seq'] for p in data['scope'])
assert [cache['W2任务'][c].value for c in ['B4','D4','F4','H4']]==[38,3,35,76]
review_ids=[new['W2 Review'].cell(r,1).value for r in range(6,30)]
assert len(review_ids)==len(set(review_ids))==24
mapped=[new['W2 Review'].cell(r,2).value for r in range(6,30) if new['W2 Review'].cell(r,2).value.startswith('UI')]
assert len(mapped)==len(set(mapped))==23
assert len(new['W2任务']._images)==7
errors=[]
for s in cache:
    for row in s:
        for c in row:
            if c.data_type=='e':errors.append([s.title,c.coordinate,c.value])
assert not errors, errors[:15]
assert new['W2用例'].tables['W2Cases'].ref=='A5:P81'
assert new['W2范围'].tables['W2Scope'].ref=='A5:S43'
for s in old:
    # 检查原输入/结果验证范围仍存在；导入器可合法拆分范围。
    ns=new[s.title]
    for d in s.data_validations.dataValidation:
        for address in str(d.sqref).split():
            area=openpyxl.worksheet.cell_range.CellRange(address)
            for coord in [openpyxl.utils.get_column_letter(area.min_col)+str(area.min_row),openpyxl.utils.get_column_letter(area.max_col)+str(area.max_row)]:
                if not any(coord in v for v in ns.data_validations.dataValidation):
                    raise AssertionError(('lost validation',s.title,coord))
with zipfile.ZipFile(path) as z:
    assert z.testzip() is None
    media=[n for n in z.namelist() if n.startswith('xl/media/')]
result={'file':str(path),'bytes':path.stat().st_size,'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'preservedCells':checked,'preservedFormulas':formula_same,'allowedModifiedCells':len(allowed),'originalSheetsRetained':len(old.sheetnames),'newSheets':new.sheetnames[len(old.sheetnames):],'weeklyCases':76,'uiCases':51,'scope':38,'formulaErrors':errors,'images':len(media),'summary':[38,3,35,76],'oldInputValidationRetained':True}
(work/'verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(result,ensure_ascii=False))
