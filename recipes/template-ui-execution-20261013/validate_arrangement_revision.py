import hashlib
import json
import re
import zipfile
from collections import Counter
from copy import copy
from html.parser import HTMLParser
from pathlib import Path

import openpyxl

root = Path('.')
work = root / '.work/template-ui-execution-20261013'
base = root / 'outputs/weekly-ui'
cfg = json.loads((work / 'arrangement-revision.json').read_text('utf8'))
old = openpyxl.load_workbook(cfg['originalWorkbook'])
new = openpyxl.load_workbook(cfg['targetWorkbook'])
assert old.sheetnames == new.sheetnames
allowed = {('业务UI与交互用例', f'{c}{r}') for r in range(2, 78) for c in ['C', 'G']}
allowed.update(('范围与依据', k) for k in cfg['scopeUpdates'])
allowed.update(('范围与依据', f"H{x['row']}") for x in cfg['scopeStatuses'])
changed = []
for before, after in zip(old, new):
    assert (before.max_row, before.max_column) == (after.max_row, after.max_column), before.title
    assert before.freeze_panes == after.freeze_panes, before.title
    assert str(before.merged_cells) == str(after.merged_cells), before.title
    assert before.data_validations == after.data_validations, before.title
    assert before.protection == after.protection
    assert before.sheet_state == after.sheet_state
    assert set(before.tables) == set(after.tables)
    for name in before.tables:
        a, b = before.tables[name], after.tables[name]
        assert a.ref == b.ref and a.autoFilter == b.autoFilter and a.tableStyleInfo == b.tableStyleInfo
    assert len(before._charts) == len(after._charts) == 0
    assert len(before._images) == len(after._images) == 0
    for arow, brow in zip(before, after):
        for a, b in zip(arow, brow):
            if a.value != b.value:
                assert (before.title, a.coordinate) in allowed, (before.title, a.coordinate)
                changed.append((before.title, a.coordinate))
            for prop in ['font', 'fill', 'border', 'alignment', 'number_format', 'protection']:
                assert copy(getattr(a, prop)) == copy(getattr(b, prop)), (before.title, a.coordinate, prop)
            assert a.hyperlink == b.hyperlink and a.comment == b.comment
            assert a.data_type == b.data_type
            assert b.data_type not in ('f', 'e')
    for key in set(before.column_dimensions) | set(after.column_dimensions):
        assert before.column_dimensions[key].width == after.column_dimensions[key].width
        assert before.column_dimensions[key].hidden == after.column_dimensions[key].hidden
    for key in set(before.row_dimensions) | set(after.row_dimensions):
        a, b = before.row_dimensions[key], after.row_dimensions[key]
        assert a.hidden == b.hidden and a.outlineLevel == b.outlineLevel
        height_allowed = (before.title == '业务UI与交互用例' and 2 <= key <= 77) or (before.title == '范围与依据' and f'A{key}' in cfg['scopeUpdates'])
        if not height_allowed:
            assert a.height == b.height, (before.title, key, 'row height')
        assert (b.height or 0) <= 409.5
assert set(changed) == allowed
assert new['业务UI与交互用例'].max_row - 1 == 76
assert Counter(new['业务UI与交互用例'].cell(r, 8).value for r in range(2, 78)) == {'未执行': 76}
assert Counter(new['通用模板UI与交互用例'].cell(r, 8).value for r in range(2, 42)) == {'通过': 12, '不通过': 10, '受阻': 18}
assert new['修改清单'].max_row - 1 == 32
assert Counter(new['修改清单'].cell(r, 2).value for r in range(2, 34)) == {'规范差异': 7, '可用性建议': 3, '用例口径修订': 1, '执行条件待补': 18, '部分步骤待补': 3}
assert Counter(new['范围与依据'].cell(r, 6).value for r in range(18, 80)) == {'UI/开': 38, '需': 24}
for update in cfg['businessUpdates']:
    assert new['业务UI与交互用例'].cell(update['row'], 3).value == update['precondition']
    assert new['业务UI与交互用例'].cell(update['row'], 7).value == update['notes']

reportname = 'SN318_通用模板40条_UI交互实测报告_20261013.html'
old_html = (work / 'before-arrangement-revision' / reportname).read_text('utf8')
html = (base / reportname).read_text('utf8')
pattern = r'<details id="UI-TPL-\d{3}" class="case".*?</details>'
assert len(re.findall(pattern, html, re.S)) == 40
assert re.findall(pattern, old_html, re.S) == re.findall(pattern, html, re.S)
assert re.findall(r'<script>.*?</script>', old_html, re.S) == re.findall(r'<script>.*?</script>', html, re.S)
assert html.count('id="current-arrangement"') == 1

class Links(HTMLParser):
    def __init__(self): super().__init__(); self.local = []
    def handle_starttag(self, tag, attrs):
        a = dict(attrs)
        for key in ('src', 'href'):
            value = a.get(key, '')
            if value and not value.startswith(('http', 'data:', '#')): self.local.append(value)

parser = Links(); parser.feed(html)
for name in parser.local: assert (base / name).is_file(), name
archive = base / 'SN318_接下来几天_UI与交互测试参考文件包_20261013.zip'
before_archive = work / 'before-arrangement-revision' / archive.name
guide = base / 'SN318_接下来几天_UI与交互测试文件说明.txt'
with zipfile.ZipFile(archive) as z, zipfile.ZipFile(before_archive) as oz:
    assert z.testzip() is None
    manifest = json.loads(z.read('文件索引.json'))
    xlsx = [n for n in z.namelist() if n.endswith('.xlsx')]
    assert len(xlsx) == 5
    primary = '用例文件/' + Path(cfg['targetWorkbook']).name
    assert z.read(primary) == Path(cfg['targetWorkbook']).read_bytes()
    assert hashlib.sha256(z.read(primary)).hexdigest() != hashlib.sha256(oz.read(primary)).hexdigest()
    unchanged = [name for name in xlsx if name != primary]
    for name in unchanged: assert z.read(name) == oz.read(name), name
    assert z.read('文件说明.txt') == guide.read_bytes()
    assert z.read('报告与证据/' + reportname) == (base / reportname).read_bytes()
    evidence = [name for name in z.namelist() if '/evidence/' in name]
    for name in evidence: assert z.read(name) == oz.read(name), name
    for entry in manifest['files']:
        data = z.read(entry['path'])
        assert len(data) == entry['bytes']
        assert hashlib.sha256(data).hexdigest() == entry['sha256']
        assert data == Path(entry['source']).read_bytes()
    assert manifest['modifiedWorkbooks'] == 1 and manifest['unchangedReferenceWorkbooks'] == 4
    assert not manifest['newApplicationTests']
    for name in parser.local: assert '报告与证据/' + name in z.namelist()
    packaged_entries = len(z.namelist())
guide_text = guide.read_text('utf8')
assert guide_text == cfg['guide']
for phrase in ['今天', '本周', '实际交付', '通用规范明确', '功能', '业务76条仍未执行']:
    assert phrase in guide_text, phrase
assert '建议安排：10/14' not in guide_text
result = {'result': 'PASS', 'changedCells': len(changed), 'authorizedChangedCellsOnly': True,
          'template40ValuesStylesEvidenceUnchanged': True, 'business76ExecutionUnchanged': True,
          'modification32ValuesStylesUnchanged': True, 'tablesValidationFreezePanesPreserved': True,
          'report40CaseDetailsAndScriptUnchanged': True, 'reportLocalLinks': len(parser.local),
          'workbooksInPackage': 5, 'unchangedReferenceWorkbooks': 4, 'evidenceFilesUnchanged': len(evidence),
          'zipEntries': packaged_entries, 'guideMatchesPackage': True, 'allPackageSha256Verified': True,
          'newPageExecution': False}
(work / 'arrangement-validation.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf8')
print(json.dumps(result, ensure_ascii=False, indent=2))
