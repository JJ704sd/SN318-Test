import json, zipfile, re, hashlib
from pathlib import Path
from collections import Counter
from copy import copy
from html.parser import HTMLParser
import openpyxl
from PIL import Image

work=Path(__file__).parent
base=Path('./artifacts/weekly-ui')
ev=base/'evidence/template-execution-20261013'
out=base/'SN318_W2_UI与交互测试用例_按工单用例风格重写.xlsx'
u=json.loads((work/'execution-update.json').read_text('utf8'))
caps=json.loads((ev/'capture-index.json').read_text('utf8'))
with zipfile.ZipFile(out) as z: assert z.testzip() is None
wb=openpyxl.load_workbook(out,data_only=False)
original=openpyxl.load_workbook(work/'workbook-before-execution.xlsx',data_only=False)
assert wb.sheetnames==original.sheetnames+['修改清单']
unchanged=wb['业务UI与交互用例'];before=original[unchanged.title]
assert list(unchanged.values)==list(before.values)
assert unchanged.freeze_panes==before.freeze_panes=='C2'
style_diffs=[]
for row in before:
    for c in row:
        n=unchanged[c.coordinate]
        for prop in ['font','fill','border','alignment','number_format','protection']:
            if copy(getattr(c,prop))!=copy(getattr(n,prop)):style_diffs.append((c.coordinate,prop))
assert not style_diffs, style_diffs[:20]
assert {k:dict(v) for k,v in unchanged.column_dimensions.items()}=={k:dict(v) for k,v in before.column_dimensions.items()}
for r,v in before.row_dimensions.items():assert unchanged.row_dimensions[r].height==v.height
for name in original.sheetnames:
    old=original[name];new=wb[name]
    assert new.freeze_panes==old.freeze_panes
    assert list(new.tables)==list(old.tables)+(['FullNormAuditIndex'] if name=='范围与依据' else [])
    assert [(str(d.sqref),d.type,d.formula1) for d in new.data_validations.dataValidation]==[(str(d.sqref),d.type,d.formula1) for d in old.data_validations.dataValidation]
    for key in old.tables:
        t=old.tables[key]
        assert new.tables[key].ref==t.ref and new.tables[key].autoFilter.ref==t.autoFilter.ref
tpl=wb['通用模板UI与交互用例'];oldtpl=original[tpl.title]
assert tpl.max_row==41 and tpl.max_column==9
assert list(next(iter(tpl.tables.values())).ref.split(':'))==['A1','I41']
maxheight=0
for row,c in enumerate(u['cases'],2):
    for j in range(1,7):
        field={1:'updatedTitle',3:'updatedPrecondition',4:'updatedSteps',5:'updatedExpected'}.get(j)
        expected=c[field] if field and c.get(field) else oldtpl.cell(row,j).value
        assert tpl.cell(row,j).value==expected,(row,j)
    assert [tpl.cell(row,j).value for j in range(7,10)]==[c['updatedNotes'],c['status'],c['actualText']]
    assert len(c['steps'])==len(c['originalSteps'])==len(c['originalExpected'])
    assert [s['step'] for s in c['steps']]==list(range(1,len(c['steps'])+1))
    maxheight=max(maxheight,tpl.row_dimensions[row].height)
    assert tpl.row_dimensions[row].height<=409.5
    for cell in tpl[row]:assert cell.alignment.wrap_text and cell.font.name=='SimSun' and cell.font.sz==11
assert Counter(tpl.cell(r,8).value for r in range(2,42))==Counter({'通过':12,'不通过':10,'受阻':18})==Counter(u['summary']['caseCounts'])
assert Counter(s['verdict'] for c in u['cases'] for s in c['steps'])==Counter({'通过':64,'不通过':11,'受阻':54})==Counter(u['summary']['stepCounts'])
assert sum(len(c['steps']) for c in u['cases'])==129
scope=wb['范围与依据'];assert scope.tables['FullNormAuditIndex'].ref=='A142:H182'
for i,a in enumerate(u['audit']['cases'],143):
    expected=[a['id'],a['before'],a['source'],a['rule']+'\n'+a['decision'],a['kind'],a['after'],a['remaining'] or '本条步骤无新增执行缺口。',a['evidence'].replace('、','\n')]
    assert [scope.cell(i,j).value for j in range(1,9)]==expected,(i,expected)
    assert scope.row_dimensions[i].height<=409.5
assert len(u['audit']['cases'])==40 and sum(a['before']=='受阻' for a in u['audit']['cases'])==20
assert sum(a['after']=='受阻' for a in u['audit']['cases'])==18
for r in [5,131]:assert '不通过10' in scope.cell(r,1).value and '受阻18' in scope.cell(r,1).value
fix=wb['修改清单'];assert fix.max_row==len(u['checklist'])+1 and fix.freeze_panes=='C2'
assert len(fix.tables)==1 and len(fix.data_validations.dataValidation)==1
for row,c in enumerate(u['checklist'],2):
    assert fix.row_dimensions[row].height<=409.5
    assert list(next(iter(fix.tables.values())).ref.split(':'))==['A1',f"K{len(u['checklist'])+1}"]
    assert [fix.cell(row,j).value for j in range(1,12)]==[c[k] for k in ['id','kind','level','object','spec','actual','recommend','accept','cases','evidence','state']]
    for proof in c['evidence'].splitlines():
        if proof.endswith('.json'):assert (ev/proof).is_file(), proof
        elif re.fullmatch(r'[ER]\d{3}',proof):assert (ev/(proof+'.jpg')).is_file(),proof
assert sum(x['kind']=='规范差异' for x in u['checklist'])==7
assert sum(x['kind']=='执行条件待补' for x in u['checklist'])==18
assert sum(x['kind']=='部分步骤待补' for x in u['checklist'])==3
assert {x for c in u['cases'] for x in c['issues']} <= {x['id'] for x in u['checklist']}
for s in wb:
    for row in s:
        for cell in row:
            assert cell.data_type not in ['f','e']
            assert not any(x in str(cell.value or '') for x in ['admin123','agent123','test.bot.example.invalid'])
assert len(caps)==102 and len({x['id'] for x in caps})==102
for cap in caps:
    path=ev/cap['file'];assert path.is_file()
    with Image.open(path) as im:assert im.format=='JPEG';im.verify()
for item in json.loads((ev/'norm-reference-manifest.json').read_text('utf8')):
    p=ev/'norm-reference'/item['name'];assert hashlib.sha256(p.read_bytes()).hexdigest()==item['sha256']
class Links(HTMLParser):
    def __init__(self):super().__init__();self.local=[];self.cases=0;self.audits=[]
    def handle_starttag(self,tag,attrs):
        a=dict(attrs)
        if tag=='details' and a.get('class')=='case':self.cases+=1
        if tag=='tr' and a.get('data-audit-id'):self.audits.append(a['data-audit-id'])
        for attr in ['src','href']:
            value=a.get(attr,'')
            if value and not value.startswith(('http','data:','#')):self.local.append(value)
p=base/'SN318_通用模板40条_UI交互实测报告_20261013.html'
parser=Links();parser.feed(p.read_text('utf8'));assert parser.cases==40
assert parser.audits==[a['id'] for a in u['audit']['cases']]
reporttext=p.read_text('utf8')
assert all(t in reporttext for t in ['通过12条','不通过10条','受阻18条','当前124张','未新增页面执行'])
for rel in parser.local:assert (base/rel).is_file(),rel
sources=json.loads((ev/'norm-full-audit-sources.json').read_text('utf8'))
assert len(sources['tables'])==sources['embeddedTableCount']==124
body=json.loads((ev/'norm-full-audit-body.json').read_text('utf8'))
assert body['revision_id']==5256 and hashlib.sha256(body['content'].encode()).hexdigest()==sources['bodySha256']
changed=[openpyxl.utils.get_column_letter(col)+str(row) for row,c in enumerate(u['cases'],2) for col,key in [(1,'updatedTitle'),(3,'updatedPrecondition'),(4,'updatedSteps'),(5,'updatedExpected')] if c.get(key) and oldtpl.cell(row,col).value!=c[key]]
report={'result':'PASS','workbook':str(out),'size':out.stat().st_size,'business76ValuesAndStylesPreserved':True,'templateUnrelatedCoreCellsPreserved':True,'updatedCaseTextCells':changed,'caseCounts':u['summary']['caseCounts'],'stepCounts':u['summary']['stepCounts'],'steps':129,'screenshots':102,'normAttachments':10,'reviewedCases':40,'reviewedOriginalBlockedCases':20,'currentReferencedTables':124,'modificationRows':len(u['checklist']),'maxTemplateRowHeightPt':maxheight,'nativeTablesAndValidationsPreserved':True,'htmlCases':40,'htmlAuditRows':40,'htmlLocalLinksValid':len(parser.local),'applicationCodeChanged':False,'newPageExecution':False}
(work/'execution-validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),'utf8')
print(json.dumps(report,ensure_ascii=False,indent=2))
