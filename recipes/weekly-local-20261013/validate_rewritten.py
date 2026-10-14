import json, re, zipfile
from pathlib import Path
from collections import Counter
import openpyxl

work = Path(__file__).parent
out = Path('./artifacts/weekly-ui/SN318_W2_UI与交互测试用例_按工单用例风格重写.xlsx')
data = json.loads((work/'rewritten-cases.json').read_text('utf-8'))
old = json.loads((work/'data.json').read_text('utf-8'))
ref = Path('D:/工单系统2.0/outputs/20260806_transfer_exploratory/工单2.0测试用例（8.4提测）-补充散发性测试用例.xlsx')
refwb = openpyxl.load_workbook(ref, read_only=True, data_only=False)
source_headers = [refwb.active.cell(1, i).value for i in range(1, 8)]
refwb.close()
with zipfile.ZipFile(out) as z:
    assert z.testzip() is None
wb = openpyxl.load_workbook(out, data_only=False)
assert wb.sheetnames == ['业务UI与交互用例', '通用模板UI与交互用例', '范围与依据']
keys = ['title','module','precondition','steps','expected','priority','notes','status','actual']
ids = []
total_steps = 0
max_height = 0
checks = []
for index, category in enumerate(['business', 'templates']):
    ws = wb.worksheets[index]
    records = data[category]
    assert ws.max_row == len(records) + 1 and ws.max_column == 9
    assert [ws.cell(1, i).value for i in range(1,8)] == source_headers
    assert [ws.cell(1, i).value for i in [8,9]] == ['执行状态','实际结果/证据']
    assert ws.freeze_panes == 'C2', ws.freeze_panes
    assert len(ws.tables) == 1
    table = next(iter(ws.tables.values()))
    assert table.ref == f'A1:I{len(records)+1}' and table.autoFilter is not None
    validations = list(ws.data_validations.dataValidation)
    assert len(validations) == 1
    dv = validations[0]
    assert str(dv.sqref) == f'H2:H{len(records)+1}' and dv.type == 'list'
    assert all(x in dv.formula1 for x in ['未执行','通过','不通过','受阻','不适用'])
    assert len(ws.conditional_formatting) >= 1
    for row, record in enumerate(records, 2):
        actual = [(ws.cell(row,col).value or '') for col in range(1,10)]
        expected = [record[k] or '' for k in keys]
        assert actual == expected, (category,row)
        ns = re.findall(r'^\[(\d+)\]', actual[3], re.M)
        ne = re.findall(r'^\[(\d+)\]', actual[4], re.M)
        assert ns == ne == list(map(str, range(1,len(ns)+1))) and len(ns) >= 3
        total_steps += len(ns)
        assert actual[7] == '未执行'
        assert record['id'] in actual[6]
        ids.append(record['id'])
        h = ws.row_dimensions[row].height
        max_height = max(max_height, h)
        assert h and h <= 409.5
        for c in ws[row]:
            assert c.alignment.wrap_text
            assert c.font.name == 'SimSun' and c.font.sz == 11
            assert c.data_type != 'f' and c.data_type != 'e'
        assert not any(term in '\n'.join(actual) for term in ['admin123', 'agent123', 'test.bot.example.invalid'])
    checks.append({'sheet':ws.title, 'cases':len(records), 'freeze':ws.freeze_panes, 'table':table.ref, 'validation':str(dv.sqref)})
assert len(ids) == len(set(ids)) == 116
assert {x['parent'] for x in data['business']} == {x['id'] for x in old['business']}
assert len({x['seq'] for x in data['business']}) == 38
assert len(data['scope']) == 62
assert sum(bool(x['new_case_ids']) for x in data['scope']) == 38
assert sum(not x['new_case_ids'] for x in data['scope']) == 24
assert sum(bool(x['actual']) for x in data['templates']) == 30
assert all('首次观察' in x['actual'] and '本轮实际结果：待填写' in x['actual'] for x in data['templates'] if x['actual'])
assert all(not x['actual'] for x in data['business'])
scope = wb['范围与依据']
assert scope.freeze_panes == 'C18' and len(scope.tables) == 2
assert len(data['reviews']) == 24
for row, review in enumerate(data['reviews'],86):
    assert scope.cell(row,1).value == review['id']
    assert scope.cell(row,4).value == review['record_id']
report = {'file':str(out), 'size_bytes':out.stat().st_size, 'cases':len(ids),'paired_steps':total_steps,'sheets':checks,'business_coverage':38,'requirement_stage_excluded':24,'historical_observations':30,'new_template_cases':8,'statuses':dict(Counter(x['status'] for x in data['business']+data['templates'])),'max_row_height_points':max_height,'result':'PASS'}
before=json.loads((work/'rewritten-cases-before-reaudit.json').read_text('utf-8'))
assert len(data['templates']) == 40
assert all(a['actual']==b['actual'] for a,b in zip(data['templates'][:30],before['templates'][:30]))
assert '提交返回403' in data['templates'][34]['title']
assert 'Toast' in data['templates'][34]['expected'] and '不把操作失败直接替换为整页403' in data['templates'][34]['expected']
assert '页面初始化' in data['templates'][38]['title']
assert all('【知识库依据】' in x['notes'] and '【审查方式】' in x['notes'] for x in data['business']+data['templates'])
assert all('【来源边界】' in x['notes'] and '【待补判据】' in x['notes'] for x in data['business'])
assert len(data['reaudit']['checks'])==8
for r,change in enumerate(data['reaudit']['checks'],115):
    assert scope.cell(r,1).value==change[0]
    assert scope.cell(r,4).value==change[2]
assert scope.cell(4,1).value.find('116')>=0
manifest=json.loads((work/'review-reaudit.manifest.json').read_text('utf-8-sig'))
assert manifest['rev']==680 and manifest['records_count']==24 and not manifest['has_more']
schedule=json.loads((work/'schedule-reaudit-revision.json').read_text('utf-8-sig'))
assert schedule['data']['revision']==988
norm=json.loads((work/'norm-reaudit-live.json').read_text('utf-8-sig'))['data']['document']
assert norm['revision_id']==5256 and '5.17. 页签 Tabs' in norm['content']
report.update(new_template_cases=10,substantive_scenarios=len(data['reaudit']['substantive_scenarios']),reaudit_entries=8,source_versions={'schedule':988,'review':680,'norm_body':5256,'norm_tables':3138})
(work/'rewrite-validation.json').write_text(json.dumps(report, ensure_ascii=False, indent=2),'utf-8')
print(json.dumps(report, ensure_ascii=False, indent=2))
