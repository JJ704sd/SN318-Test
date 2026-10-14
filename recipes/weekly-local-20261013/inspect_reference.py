import json
from pathlib import Path
import openpyxl

p=Path('D:/工单系统2.0/outputs/20260806_transfer_exploratory/工单2.0测试用例（8.4提测）-补充散发性测试用例.xlsx')
w=openpyxl.load_workbook(p,data_only=False)
out=[]
for s in w:
    rows=[]
    for cells in s.iter_rows():
        vals=[c.value for c in cells]
        if any(v is not None for v in vals):rows.append({'row':cells[0].row,'values':vals})
    cols={k:{'width':v.width,'hidden':v.hidden} for k,v in s.column_dimensions.items()}
    info={'name':s.title,'max_row':s.max_row,'max_col':s.max_column,'merged':[str(x) for x in s.merged_cells.ranges],'freeze':str(s.freeze_panes),'filter':s.auto_filter.ref,'columns':cols,'rows':rows}
    out.append(info)
    print(json.dumps({k:info[k] for k in ['name','max_row','max_col','merged','freeze','filter','columns']},ensure_ascii=False,default=str))
    sample=rows if '散发' in s.title or '补充' in s.title else rows[:10]+rows[-3:]
    for r in sample[:35]:print(json.dumps(r,ensure_ascii=False,default=str))
Path(__file__).with_name('reference-analysis.json').write_text(json.dumps(out,ensure_ascii=False,indent=2,default=str),'utf-8')
