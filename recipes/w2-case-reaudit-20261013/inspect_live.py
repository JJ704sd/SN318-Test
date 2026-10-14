import json,re,difflib,hashlib
from pathlib import Path
from html import unescape
import xml.etree.ElementTree as ET
import openpyxl
W=Path(__file__).parent
def load(p):return json.loads(Path(p).read_text(encoding='utf-8-sig'))
def doc(f):return load(W/f)['data']['document']
p=doc('product-live.json');n=doc('norm-live.json')
print('VERSIONS',{'schedule':load(W/'schedule-info.json')['data']['revision'],'product':p['revision_id'],'norm':n['revision_id'],'review':load(W/'review-live.manifest.json')['rev']})
old=load('./recipes/w2-week-20261012/product-live.json')['data']['document']
changes=[]
aa=re.findall('<[^>]+>|[^<]+',old['content']);bb=re.findall('<[^>]+>|[^<]+',p['content'])
for op,i,j,k,l in difflib.SequenceMatcher(None,aa,bb).get_opcodes():
 if op!='equal':changes.append({'old':''.join(aa[i:j]),'new':''.join(bb[k:l])})
(W/'product-changes.json').write_text(json.dumps(changes,ensure_ascii=False,indent=2),encoding='utf-8')
print('CHANGES',json.dumps(changes,ensure_ascii=False)[:6000])
root=ET.fromstring('<root>'+p['content']+'</root>')
blocks=[]
for e in root.iter():
 if e.tag in ['p','h1','h2','h3','h4','li','title'] and e.get('id'):
  text=' '.join(''.join(e.itertext()).split())
  if text:blocks.append({'id':e.get('id'),'tag':e.tag,'text':text,'comments':e.get('comment-refs')})
(W/'product-blocks.json').write_text(json.dumps(blocks,ensure_ascii=False,indent=2),encoding='utf-8')
print('HEADINGS',[(b['id'],b['text']) for b in blocks if b['tag'] in ['h1','h2','h3','h4']])
print('COMMENTS', {k:v for k,v in p.get('reference_map',{}).get('comments',{}).items() if k in ['c18','c19','c20','c21','c22','c23']})
path=Path('./artifacts/w2-tasks/20261013/PCS_W2前端测试用例_按任务节奏细化.xlsx')
w=openpyxl.load_workbook(path)
rows=[]
for r in range(6,103):rows.append([w['前端用例'].cell(r,c).value for c in range(1,18)])
(W/'input-readback.json').write_text(json.dumps({'path':str(path),'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'cases':rows,'sheets':w.sheetnames},ensure_ascii=False,indent=2),encoding='utf-8')
print('RECORDS',len(rows),'execution inputs',[(r[0],r[10:15]) for r in rows if r[10]!='待评审' or r[11] not in ['未执行','不适用'] or r[12] not in ['未执行','不适用'] or r[13] or r[14]])
print('REVIEW',[(r['record_id'],r.get('页面/模块'),r.get('框架适配判断'),r.get('确认状态'),r.get('Review 建议')) for r in [json.loads(x) for x in (W/'review-live.ndjson').read_text(encoding='utf-8-sig').splitlines()] if r['record_id'] in ['review-67fb39add5','review-4122ab1853','review-294de09d51','review-a7e8b941e2','review-24220c26b9','review-331da30eb4','review-f828afea0f']])
