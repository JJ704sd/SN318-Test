import json, hashlib, re
from pathlib import Path
from zipfile import ZipFile
import openpyxl
from copy import copy
W=Path(__file__).parent
d=json.loads((W/'revision-data.json').read_text(encoding='utf-8'))
input_path=Path(d['input']); output_path=Path(d['out'])
orig=openpyxl.load_workbook(input_path); out=openpyxl.load_workbook(output_path)
cached=openpyxl.load_workbook(output_path,data_only=True)
assert hashlib.sha256(input_path.read_bytes()).hexdigest()==d['input_hash']
assert out.sheetnames==orig.sheetnames+['复核记录']
s=out['前端用例']; end=110
assert [s.cell(r,1).value for r in range(6,111)]==[r[0] for r in d['cases']]
for row in range(6,103):
    for col in range(11,18):
        assert s.cell(row,col).value==orig['前端用例'].cell(row,col).value,(row,col)
assert [cached['本周执行'][a].value for a in ['B4','D4','F4','H4']]==[105,38,105,0]
assert all(cached['前端用例'].cell(r,17).value=='待评审/依据' for r in range(6,111))
assert sum(s.cell(r,1).value.startswith('RV') for r in range(6,111))==14
assert all(s.cell(r,12).value=='不适用' and s.cell(r,13).value=='不适用' for r in [103,104])
assert all(s.cell(r,12).value=='未执行' and s.cell(r,13).value=='未执行' for r in range(105,111))
assert len(out['工作包对应'].tables)==1 and out['工作包对应'].tables['PackageMapping'].ref=='A5:P43'
assert s.tables['FrontendCases'].ref=='A5:S110'
assert out['本周执行'].tables['CaseSources'].ref=='A34:D41'
assert out['复核记录'].tables['CaseReaudit'].ref=='A5:G110'
assert out['复核记录'].tables['PageEvidenceIndex'].ref=='A114:E134'
assert len(out['复核记录']._images)==20
assert s.freeze_panes==orig['前端用例'].freeze_panes
assert out['工作包对应'].freeze_panes==orig['工作包对应'].freeze_panes
dv=list(s.data_validations.dataValidation)
assert any('K103:K110' in str(x.sqref) for x in dv)
assert any('L103:M110' in str(x.sqref) for x in dv)
assert len(s.conditional_formatting)>=len(orig['前端用例'].conditional_formatting)
errors=[]
for sh in cached:
    for row in sh:
        for c in row:
            if c.data_type=='e':errors.append((sh.title,c.coordinate,c.value))
assert not errors,errors
# 原有非目标区域与基础格式保留；行高随新文字长度调整。
for row in range(27,32):
    for col in range(1,4): assert out['本周执行'].cell(row,col).value==orig['本周执行'].cell(row,col).value
for row in range(6,103):
    for col in range(11,17):
        assert s.cell(row,col).number_format==orig['前端用例'].cell(row,col).number_format
        assert copy(s.cell(row,col).fill)==copy(orig['前端用例'].cell(row,col).fill)
# 38包映射双向一致：公共条目独立复用，FE22条件能力不伪造固定包。
public=set(d['public']);ids={r[0] for r in d['cases']}
for p in d['packages']:
    seq=p[0]
    for column,prefix in [(9,'RV'),(10,'FE')]:
        mapped=p[column].split('、') if p[column] else []
        expected=[id for id in ids if id.startswith(prefix) and id not in public and seq in d['scopes'][id]]
        assert set(mapped)==set(expected),(seq,column)
        assert len(mapped)==len(set(mapped))
    assert p[11]==f'BP{seq:02d}'
    rr=next(r for r in range(6,44) if out['工作包对应'].cell(r,1).value==seq)
    assert [out['工作包对应'].cell(rr,c).value or '' for c in range(10,17)]==[x or '' for x in p[9:16]]
assert all(r[0] in ids for r in d['journal'])
assert len({r[0] for r in d['journal']})==105
with ZipFile(output_path) as z:
    assert len([x for x in z.namelist() if x.startswith('xl/media/')])==20
result={'output':str(output_path),'sha256':hashlib.sha256(output_path.read_bytes()).hexdigest(),'bytes':output_path.stat().st_size,'stats':d['stats'],'sourceUnchanged':True,'originalKtoQPreserved':True,'bidirectionalPackageMappingValid':True,'embeddedImages':20,'nativeTables':{sh.title:len(sh.tables) for sh in out},'cachedSummary':[105,38,105,0],'formulaErrors':errors}
(W/'verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(result,ensure_ascii=False))
