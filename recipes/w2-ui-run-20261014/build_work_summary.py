from pathlib import Path
from collections import Counter
from datetime import datetime, timezone, timedelta
import json, hashlib, zipfile, xml.etree.ElementTree as ET, posixpath, re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path('.')
WORK=ROOT/'.work/w2-ui-run-20261014'
OUT=ROOT/'outputs/ui-audit-20261014/20261014-ui-audit'
OLD=ROOT/'outputs/template-reaudit/受阻16条核实/SN318_W2_UI与交互测试用例_16条受阻核实_含人工二次复核.xlsx'
NEW=OUT/'SN318_W2_UI与交互_20261014_dev_b2efcdb_dirty-5c6bc530.xlsx'
DEST=OUT/'SN318_前端UI规范审查与修改清单_两日工作概述_20261014.docx'
record=json.loads((OUT/'审查记录.json').read_text(encoding='utf-8'))
issues={i['id']:i for i in record['issues']}
hash_before={str(p):hashlib.sha256(p.read_bytes()).hexdigest() for p in [OLD,NEW]}

def read_xlsx(p):
 ns='http://schemas.openxmlformats.org/spreadsheetml/2006/main';Q=lambda t:'{'+ns+'}'+t
 with zipfile.ZipFile(p) as z:
  strings=[''.join(e.itertext()) for e in ET.fromstring(z.read('xl/sharedStrings.xml'))] if 'xl/sharedStrings.xml' in z.namelist() else []
  rels={e.attrib['Id']:e.attrib['Target'] for e in ET.fromstring(z.read('xl/_rels/workbook.xml.rels'))}
  result={}
  for sheet in ET.fromstring(z.read('xl/workbook.xml')).find(Q('sheets')):
   rel=rels[sheet.attrib['{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id']]
   cells={}
   for cell in ET.fromstring(z.read(rel.lstrip('/') if rel.startswith('/') else posixpath.normpath('xl/'+rel))).iter(Q('c')):
    v=cell.find(Q('v'));val=v.text if v is not None else ''
    if cell.attrib.get('t')=='s':val=strings[int(val)]
    elif cell.attrib.get('t')=='inlineStr':val=''.join(cell.find(Q('is')).itertext())
    cells[cell.attrib['r']]=val
   result[sheet.attrib['name']]=cells
  return result
old,new=read_xlsx(OLD),read_xlsx(NEW)
assert all(new[s].get(a,'')==v for s,cells in old.items() for a,v in cells.items())
tpl=Counter(new['通用模板UI与交互用例']['K'+str(r)] for r in range(2,42))
bus=Counter(new['业务UI与交互用例']['K'+str(r)] for r in range(78,103))
assert dict(tpl)==record['counts']['template'] and dict(bus)==record['counts']['business']
assert len(issues)==16
assert new['通用模板UI与交互用例']['J2']==old['通用模板UI与交互用例']['J2']=='不通过'
assert len(record['packages'])==38

doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=Inches(.70);sec.bottom_margin=Inches(.68)
sec.left_margin=Inches(.76);sec.right_margin=Inches(.76)
sec.header_distance=Inches(.25);sec.footer_distance=Inches(.28)
FONT='Microsoft YaHei'
def font_style(style,size,bold=False):
 style.font.name=FONT;style.font.size=Pt(size);style.font.bold=bold;style.font.color.rgb=RGBColor(0,0,0)
 for color in style.element.findall('.//'+qn('w:color')):
  for attr in ['themeColor','themeTint','themeShade']:color.attrib.pop(qn('w:'+attr),None)
 fonts=style.element.get_or_add_rPr().get_or_add_rFonts()
 for prop in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+prop),FONT)
 for key in list(fonts.attrib):
  if 'theme' in key.lower():fonts.attrib.pop(key,None)
 pr=style.element.get_or_add_rPr()
 for tag in ['spacing','kern']:
  for item in pr.findall(qn('w:'+tag)):pr.remove(item)
 size_cs=pr.find(qn('w:szCs'))
 if size_cs is None:size_cs=OxmlElement('w:szCs');pr.append(size_cs)
 size_cs.set(qn('w:val'),str(int(size*2)))
 for b in style.element.findall('.//'+qn('w:pBdr')):b.getparent().remove(b)
font_style(doc.styles['Normal'],11)
normal=doc.styles['Normal'].paragraph_format
normal.line_spacing=1.15;normal.space_after=Pt(4);normal.widow_control=True
font_style(doc.styles['Title'],21,True)
doc.styles['Title'].paragraph_format.space_after=Pt(10)
for name,size in [('Heading 1',15),('Heading 2',12)]:
 font_style(doc.styles[name],size,True)
 doc.styles[name].paragraph_format.space_before=Pt(9)
 doc.styles[name].paragraph_format.space_after=Pt(4)
 doc.styles[name].paragraph_format.keep_with_next=True
for style in ['Header','Footer']:font_style(doc.styles[style],9)
doc.core_properties.title='SN318 前端UI规范审查与修改清单'
doc.core_properties.subject='2026年10月13日至14日工作概述及W2整改交接'
doc.core_properties.author='SN318'
doc.core_properties.keywords='SN318 W2 UI规范审查 修改清单 验收标准'
doc.core_properties.comments='依据既有工作簿与留存证据整理，未新增页面执行。'

def p(text='',bold_prefix=None,small=False,keep=False):
 para=doc.add_paragraph();para.paragraph_format.keep_with_next=keep
 if bold_prefix and text.startswith(bold_prefix):
  para.add_run(bold_prefix).bold=True;para.add_run(text[len(bold_prefix):])
 else:para.add_run(text)
 if small:
  for run in para.runs:run.font.size=Pt(9.5)
  para.paragraph_format.line_spacing=1.1;para.paragraph_format.space_after=Pt(4)
 return para
next_page=False
def heading(text,level=1):
 global next_page
 para=doc.add_paragraph(text,'Heading '+str(level))
 if next_page:para.paragraph_format.page_break_before=True;next_page=False
 return para
def page():
 global next_page
 next_page=True
def hyperlink(para,label,target,small=False):
 link=OxmlElement('w:hyperlink');link.set(qn('r:id'),para.part.relate_to(target,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True))
 run=OxmlElement('w:r');pr=OxmlElement('w:rPr');fonts=OxmlElement('w:rFonts')
 for prop in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+prop),FONT)
 pr.append(fonts);c=OxmlElement('w:color');c.set(qn('w:val'),'046EFC');pr.append(c)
 size=OxmlElement('w:sz');size.set(qn('w:val'),'19' if small else '22');pr.append(size)
 run.append(pr);t=OxmlElement('w:t');t.text=label;run.append(t);link.append(run);para._p.append(link)
def bookmark(para,name):
 id=str(len(doc.element.findall('.//'+qn('w:bookmarkStart')))+1)
 start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),id);start.set(qn('w:name'),name)
 end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),id)
 para._p.insert(0,start);para._p.append(end)
def refs(item):
 cases=item['case_ids']
 if cases==[f'UI-M1-{i:03d}' for i in range(1,21)]:case_text='UI-M1-001至020'
 elif len(cases)>8 and all(c.startswith('UI-TPL-') for c in cases):case_text='UI-TPL-'+ '/'.join(c.split('-')[-1] for c in cases)
 else:case_text='、'.join(cases)
 para=p('关联用例 '+case_text+'。证据 ',small=True)
 for index,e in enumerate(item['evidence_ids']):
  path=OUT/'evidence'/f'{e}.json';assert path.is_file()
  if index:para.add_run('；')
  hyperlink(para,e,'evidence/'+e+'.json',True)
 return para
def table(headers,rows,widths):
 t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for c,w in zip(t.columns,widths):c.width=Inches(w)
 for row,values in [(t.rows[0],headers)]+[(t.add_row(),r) for r in rows]:
  for j,(cell,text) in enumerate(zip(row.cells,values)):
   cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   para=cell.paragraphs[0];para.paragraph_format.space_after=Pt(2);para.paragraph_format.space_before=Pt(2);para.paragraph_format.line_spacing=1.12
   run=para.add_run(str(text));run.font.size=Pt(10.5)
   if j==0 or all(len(str(x))<12 for x in values):para.alignment=WD_ALIGN_PARAGRAPH.CENTER
   tcPr=cell._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
   for side in ['top','bottom','left','right']:
    e=OxmlElement('w:'+side);e.set(qn('w:w'),'85' if side in ['top','bottom'] else '110');e.set(qn('w:type'),'dxa');m.append(e)
   tcPr.append(m)
   borders=OxmlElement('w:tcBorders')
   for side in ['top','left','bottom','right']:
    e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
   tcPr.append(borders)
  prevent=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(prevent)
 for cell in t.rows[0].cells:
  shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DCE6F1');cell._tc.get_or_add_tcPr().append(shade)
  for run in cell.paragraphs[0].runs:run.bold=True;run.font.color.rgb=RGBColor(0,0,0)
 repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
 p('').paragraph_format.space_after=Pt(1)
 return t

header=sec.header.paragraphs[0];header.text='SN318 W2 UI审查工作清单';header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER;footer.add_run('第 ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field);footer.add_run(' 页')

# The first page answers the manager's two-day progress question before technical detail.
doc.add_paragraph('SN318 前端UI规范审查与修改清单','Title')
p('2026年10月13日至14日工作概述  W2阶段交付',small=True)
p('这两天围绕SN318本地前端开展了通用模板用例复核、已交付基础配置页面的UI规范审查和交互补测，形成可交接的修改建议与验收标准。当前整理出11项已确认差异、3项待确认事项和2组执行缺口，供工作汇报、前端整改和后续复测使用。')
heading('这两天完成的工作')
table(['时间','完成事项','形成的记录'],[
 ['10/13','对照通用设计规范检查模板的视觉、交互与边界状态，梳理W2用例及可测范围。','保留模板用例、业务候选框架、实际结果和截图等证据；缺入口或判据的项单独记录。'],
 ['10/14','复核40条模板用例，重点核实原16条受阻项；再对已交付M1页面开展查询、重置、关联页签和定向表单检查。','原受阻核实稿保留；新增25条业务补测，输出16项当前修改或补齐记录，并保留人工二次复核。']
],[.73,3.12,3.13])
heading('本次审查结果')
table(['范围','通过','不通过','受阻'],[
 ['模板40条',tpl['通过'],tpl['不通过'],tpl['受阻']],
 ['新增业务补测25条',bus['通过'],bus['不通过'],bus['受阻']]
],[3.65,1.11,1.11,1.11])
p('新增业务补测覆盖18个M1列表、2个关联子表和5个定向检查。20条空结果用例关联同一个共性问题R09，用例不通过数量不能直接当作独立缺陷数量。原76条W2业务候选仍未完整执行。',small=True)
p('本批优先推进Tablet表单双列、查询加载反馈和筛选无结果引导等P1问题。修改清单已经形成，前端整改与整改复测尚未完成；本周继续随实际页面和修复Build交付扩大覆盖。')

titles={
 'W2-R01':'按钮圆角与规范不一致','W2-R02':'标准页头高度超出规范','W2-R03':'操作列对齐与按钮间距不一致','W2-R04':'普通正文使用了标题颜色','W2-R05':'运单链接未使用规定品牌色','W2-R06':'Tablet表单仍显示双列','W2-R07':'查询期间按钮没有加载与禁用反馈','W2-R08':'刷新列表后退出当前登录状态','W2-R09':'筛选无结果与首次无数据引导混用','W2-R10':'上传说明与实际触发行为不一致','W2-R11':'汇率数字没有右对齐'
}
def add_issue(id):
 start=len(doc.paragraphs)
 i=issues[id];h=heading(id.replace('W2-','')+' '+titles[id],2);bookmark(h,id.replace('-','_'))
 p('问题ID '+id+'  '+i['level']+'  '+i['type']+'  对象 '+i['page'],small=True,keep=True)
 p('依据与差异：'+i['norm']+'。'+i['actual']+'。',bold_prefix='依据与差异：')
 p('复现与修改：'+i['repro']+'。'+i['recommendation']+'。',bold_prefix='复现与修改：')
 acceptance=i['acceptance']
 if id=='W2-R06':acceptance='768px≤视口宽度<1280px时采用单列，标签在上、两侧padding16px；底栏键盘可达、无页面横向溢出。1280px起恢复Desktop布局。'
 if id=='W2-R08':
  acceptance='先确认本地Mock是否要求刷新保持会话。采纳原用例时，刷新仍停留/demo/list，导航和详情可用；若采用不持久登录的演示方案，需明确修订前置与验收口径，并保留差异记录。'
 if id=='W2-R10':
  acceptance='按确认的上传方案验收。若采用页面声明的显式提交，选择后仅待上传，明确开始才显示进度；取消能终止，失败重试后可恢复完成。本项只验证Mock界面行为，未证明真实文件传输。'
 p('验收标准：'+acceptance, bold_prefix='验收标准：')
 if i['needed']:p('剩余条件：'+i['needed']+'。',small=True)
 refs(i)
 if id=='W2-R09':p('空结果完成态以DOM记录为依据；部分截图仍为旧列表画面，不能作为同一时刻的空态截图。',small=True)
 for para in doc.paragraphs[start:-1]:para.paragraph_format.keep_with_next=True

page();heading('优先整改的交互与布局问题')
p('以下11项均为当前已确认差异，状态为待整改且未完成整改复测，责任人待分派。问题等级沿用当前清单；工作包的P0或P1排期优先级另行记录。',small=True)
for id in ['W2-R06','W2-R07','W2-R09','W2-R08']:add_issue(id)
heading('公共组件视觉差异')
for id in ['W2-R01','W2-R02','W2-R03','W2-R04']:add_issue(id)
heading('链接上传与数值呈现差异')
for id in ['W2-R05','W2-R10','W2-R11']:add_issue(id)
p('复测应检查修复项及相邻状态，例如默认、悬浮、加载、禁用和错误恢复。公共问题可以关联多个页面，但每个业务页面的字段、版本和权限呈现仍需分别验证。')

heading('需要确认和补齐的事项')
p('以下3项需明确产品或设计口径，2组记录执行条件缺口。责任人均待确认；这些项不能直接计作已证实产品缺陷。',small=True)
for id,title in [('W2-D01','费用项字段与字典来源'),('W2-D02','汇率版本维护与审批规则'),('W2-D03','Mock数据来源的界面说明')]:
 start=len(doc.paragraphs)
 i=issues[id];heading(id.replace('W2-','')+' '+title,2)
 p('问题ID '+id+'  '+i['level']+'  '+i['type']+'  对象 '+i['page'],small=True,keep=True)
 p('现状与依据：'+i['norm']+'。'+i['actual']+'。',bold_prefix='现状与依据：')
 acceptance='产品确认统一口径并记录决定，补充版本变更、权限呈现和提交异常样例，再按确认口径复测' if id=='W2-D02' else i['acceptance']
 p('需要确认：'+i['recommendation']+'。确认后的验收要求：'+acceptance+'。',bold_prefix='需要确认：')
 p('当前状态：'+i['status']+'；具体负责人待确认。',small=True);refs(i)
 for para in doc.paragraphs[start:-1]:para.paragraph_format.keep_with_next=True
start=len(doc.paragraphs)
i=issues['W2-G01'];heading('G01 异常响应和恢复样例',2)
p('问题ID W2-G01  P1  样例缺口  当前状态 受阻待补齐',small=True,keep=True)
p('需提供可从页面触发的timeout、offline、500、提交403、初始化403及恢复成功样例，记录输入保留、局部错误和按钮解锁。直接访问/403或/500展示路由仅能证明展示，不替代接口异常链路。')
p('验收标准：'+i['acceptance']+'。当前状态：受阻待补齐；响应与样例入口需要开发和测试协同，具体负责人待确认。',small=True);refs(i)
for para in doc.paragraphs[start:-1]:para.paragraph_format.keep_with_next=True
start=len(doc.paragraphs)
heading('G02 专项判据样例与本轮未完成步骤',2)
p('问题ID W2-G02  P2  执行条件或取证缺口  当前状态 待补齐或继续执行',small=True,keep=True)
p('需确认Mobile范围、未保存离开与草稿规则、跨页选择及可访问性验收策略；补齐Tabs数量与长标签、长文本、加载完成和重试成功等专项样例。')
p('本轮未完成的横向滚动、抽屉关闭路径、部分组件状态及同状态示例图对照，应由测试继续操作与取证；不能统一归因于开发缺样例。具体子步骤以当前工作簿L列为准。',small=True)
p('验收标准：逐项具备明确预期和可执行条件，完成操作并记录结果及证据后再更新状态；责任人待确认。',small=True)
refs(issues['W2-G02'])
for para in doc.paragraphs[start:-1]:para.paragraph_format.keep_with_next=True

heading('本周推进安排与覆盖边界')
p('W2为2026年10月12日至16日，排期38项处于UI与开发并行阶段，其中P0有24项、P1有14项。按实际提测入口开展全面介入；本批完整工作包通过数为0。')
table(['W2范围','本批覆盖','后续动作'],[
 ['基础配置9项','费用项、币种与汇率、业务字典3项有M1部分证据，其余6项本轮未取得对应入口。','确认M1与专项PRD映射，补字段、版本、权限及异常分支。'],
 ['服务管理10项','本轮未取得对应业务入口。','收到可测页面与样例后，按P0优先核对列表、表单、详情及引用关系。'],
 ['价格维护18项','地点和Zone的主数据补测不替代价卡、分区、导入等专项页面。','随专项页面提测核对计费明细、分区选择、导入与边界状态。'],
 ['系统配置与基础数据1项','操作日志本轮未取得对应专项入口。','确认日志筛选、只读详情和异常样例后执行。']
],[1.50,2.71,2.77])
heading('后续工作触发条件')
p('页面提测时核对入口、Build、角色、设计和样例，先查通用规范与主要操作，再补专项字段与边界；缺判据只阻塞依赖该判据的步骤。')
p('修复Build交付后优先复测P1交互及公共组件问题。周内收尾汇总已执行页面、关闭项、剩余受阻和下一批入口；提测日与负责人按实际交付确认。')
heading('记录口径与证据使用')
p('审查基线为dev b2efcdb，dirty指纹5c6bc530，使用Mock认证和数据，仅验该快照的前端表现；真实接口、权限及业务持久化未验。',small=True)
p('本次结果看模板K至M列、业务78至102行、清单35行起、汇总184行起。原执行记录及人工意见保留，历史与本轮分别计数。',small=True)
p('截图与DOM顺序采样；未同步时引用对应量测并注明限制。人工二次复核沿用原工作簿，本文件不代填人工结论。',small=True)
start=len(doc.paragraphs)
heading('交付依据与附件')
para=p('工作底稿：',small=True);hyperlink(para,'当前执行稿与清单',NEW.name,True);para.add_run('；');hyperlink(para,'历史受阻核实稿',OLD.as_uri(),True)
para=p('证据材料：',small=True);hyperlink(para,'审查报告','审查报告.html',True);para.add_run('；');hyperlink(para,'完整证据包','SN318_W2_UI审查_20261014_证据包.zip',True)
para=p('规范依据：',small=True);hyperlink(para,'通用设计规范V1.2','https://example.invalid/source/05cc32fecc',True)
p('规范正文rev5256及本次涉及的嵌入表rev3138；排期rev988；产品PRD rev8350、基础PRD rev2。相关快照随证据资料保存。本次仅将既有审查成果整理为工作清单，未新增页面实测或源码修改。',small=True)
for para in doc.paragraphs[start:-1]:para.paragraph_format.keep_with_next=True

doc.save(DEST)
assert all(hashlib.sha256(Path(path).read_bytes()).hexdigest()==value for path,value in hash_before.items())
with zipfile.ZipFile(DEST) as z:
 assert z.testzip() is None
 xml=z.read('word/document.xml').decode('utf-8')
 assert all(id in xml for id in issues)
 assert '不通过' in xml
 headings=[p.text for p in doc.paragraphs if p.style.name in ['Title','Heading 1','Heading 2']]
 assert all(not re.search(r'[，。；：、（）()\-]',h) for h in headings)
check={'output':str(DEST),'source_sha256':hash_before,'original_records_unchanged':True,'template_results':dict(tpl),'business_results':dict(bus),'issue_ids':list(issues),'work_package_count':38,'human_review':'source value UI-TPL-001 retained; no AI human review','repackaging_only':True,'planned_pages':6,'local_evidence_links_exist':True,'created_at':datetime.now(timezone(timedelta(hours=8))).isoformat(timespec='seconds')}
(WORK/'document-content-validation.json').write_text(json.dumps(check,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(check,ensure_ascii=False))
