import json,re
from pathlib import Path
from collections import Counter
import openpyxl
root=Path('./recipes/sn318-site-review')
folder=Path('./artifacts/smoke-ui')
src=folder/'SN318_功能冒烟与界面交互验证_策略调整.xlsx'
dst=folder/'SN318_功能冒烟与界面交互验证_覆盖复核.xlsx'
before=openpyxl.load_workbook(src,data_only=False)
w=openpyxl.load_workbook(dst,data_only=False)
c=openpyxl.load_workbook(dst,data_only=True)
assert w.sheetnames==before.sheetnames
def cells(sh):
    return {x.coordinate:x.value for row in sh for x in row if x.value is not None}
unchanged=['冒烟用例','需求对照','网站详情对照','UI修改清单']
for name in unchanged:
    assert cells(w[name])==cells(before[name]),name
preserved=[]
for name,areas in {'本批提测':['B7:B17','D40:G61','D115:G117','D156:G160','B33:B36'],'UI检查':['D6:H54','M6:O54','U6:V54'],'开发冒烟':['G7:M22'],'UI走查':['H7:R55']}.items():
    for area in areas:
        assert [[x.value for x in row] for row in w[name][area]]==[[x.value for x in row] for row in before[name][area]],(name,area)
        preserved.append([name,area])
for name,last,prefix in [('冒烟用例',96,'SM'),('UI检查',54,'UI')]:
    assert [w[name].cell(r,1).value for r in range(6,last+1)]==[f'{prefix}{r-5:02d}' for r in range(6,last+1)]
    assert all(w[name].cell(r,7).value=='未执行' for r in range(6,last+1))
assert Counter(w['冒烟用例'].cell(r,16).value for r in range(6,97))=={'主流程回归':28,'专项回归':63}
for n in range(1,50):
    fr=n+5;wr=n+6
    assert w['UI走查'][f'A{wr}'].value==f'UI{n:02d}'
    assert w['UI走查'][f'I{wr}'].value=='未走查'
    assert w['UI走查'][f'U{wr}'].value=='待确认'
    for col,ref in [('F','J'),('G','K'),('T','I')]:
        assert w['UI走查'][f'{col}{wr}'].value==f"='UI检查'!{ref}{fr}"
    assert '页面覆盖未通过' in w['UI检查'][f'I{fr}'].value
    assert f"'UI走查'!Y{wr}" in w['UI检查'][f'I{fr}'].value
    assert c['UI走查'][f'V{wr}'].value==c['UI走查'][f'W{wr}'].value==c['UI走查'][f'X{wr}'].value==0
    assert c['UI走查'][f'Y{wr}'].value=='页面范围待核定'
assert len({w['UI走查'][f'A{r}'].value for r in range(62,72)})==10
for r in range(62,112):
    assert w['UI走查'][f'E{r}'].value=='待选择'
    assert w['UI走查'][f'H{r}'].value=='未执行'
    for col in ['O','P','Q','R']:assert w['UI走查'][f'{col}{r}'].data_type=='f'
    assert c['UI走查'][f'O{r}'].value==('未纳入' if r<72 else None)
assert '对象专属规则待确认' in w['开发冒烟']['N9'].value
assert '$B$11' not in w['开发冒烟']['N9'].value
assert w['开发冒烟']['Q9'].value is None and w['开发冒烟']['R9'].value is None
assert 'DS14' in w['开发冒烟']['D9'].value and 'DS15' in w['开发冒烟']['D9'].value
for n in range(1,17):assert w['开发冒烟'][f'I{n+6}'].value=='未执行'
for a in ['B22','F22','H22']:assert c['本批提测'][a].value==0
assert c['开发冒烟']['E4'].value==0
assert c['UI走查']['E4'].value==0
assert c['UI修改清单']['B4'].value==21
assert c['UI修改清单']['E4'].value==c['UI修改清单']['H4'].value==0
assert c['本批提测']['B6'].value=='未就绪'
errors=[(sh.title,x.coordinate,x.value) for sh in c for row in sh for x in row if x.data_type=='e']
assert not errors,errors
for sh in w:
    assert not any('复核模拟' in str(x.value) or x.value=='AUDIT-BUILD' or x.value=='REQ-AUDIT' or x.value=='UI-AUDIT' or x.value in ['AUDIT-P1','AUDIT-P2'] for row in sh for x in row if x.data_type!='f'),sh.title
refs={sh:{t.name:t.ref for t in w[sh].tables.values()} for sh in w.sheetnames}
assert refs['UI走查']['UIWalkthrough']=='A6:Y55'
assert refs['UI走查']['UIPageExecution']=='A61:R111'
assert refs['开发冒烟']['DeveloperSmoke']=='A6:R22'
for name in unchanged:
    assert refs[name]=={t.name:t.ref for t in before[name].tables.values()}
live=[json.loads(l) for l in (root/'ui-final-review-live.ndjson').read_text(encoding='utf-8-sig').splitlines()]
valid={r['record_id']:r for r in live if any(r.get(k) for k in ['页面/模块','当前原型问题','Review 建议','后续处理建议'])}
mapped={w['UI检查'].cell(r,28).value:r for r in range(6,55) if w['UI检查'].cell(r,28).value}
assert set(mapped)==set(valid) and len(mapped)==21
for rid,r in mapped.items():
    assert w['UI检查'].cell(r,17).value=='、'.join(valid[rid]['确认状态'])
    assert w['UI检查'].cell(r,18).value==re.sub(r'^[^\u4e00-\u9fff]+','',valid[rid]['框架适配判断'][0])
case_map={f'{prefix}{r-5:02d}':(name,r,col) for name,last,prefix,col in [('冒烟用例',96,'SM','V'),('UI检查',54,'UI','AF')] for r in range(6,last+1)}
declared={ident:set() for ident in case_map}
for dr in list(range(40,62))+list(range(115,118))+list(range(156,161)):
    for m in re.finditer(r'(SM|UI)(\d+)(?:[—–-](?:SM|UI)?(\d+))?',str(w['本批提测'].cell(dr,3).value)):
        for n in range(int(m[2]),int(m[3] or m[2])+1):declared[f'{m[1]}{n:02d}'].add(dr)
for ident,(sh,r,col) in case_map.items():
    actual=set(map(int,re.findall(r"'本批提测'!\$H\$(\d+)",w[sh][f'{col}{r}'].value)))
    assert actual==declared[ident],ident
report={'path':str(dst),'bytes':dst.stat().st_size,'reviewRevision':642,'sheetsPreserved':True,'formalCases':140,'developerScenes':16,'pageCandidateExamples':10,'pageExecutionInputRows':50,'formalPasses':0,'allUnexecuted':True,'formulaErrors':0,'exactDecisionEdges':True,'inputsPreserved':preserved,'unchangedSheets':unchanged,'tableRefs':refs}
(root/'coverage-saved-verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(report,ensure_ascii=False))
