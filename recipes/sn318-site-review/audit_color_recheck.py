import json,re
from pathlib import Path
import openpyxl
root=Path('./recipes/sn318-site-review')
folder=Path('./artifacts/smoke-ui')
w=openpyxl.load_workbook(folder/'SN318_提测前冒烟与UI规范_颜色层级对应.xlsx',data_only=False)
b=w['本批提测']
decision_rows=list(range(40,62))+list(range(115,118))+list(range(156,161))
def ids(v):
    return [f'{m[1]}{n:02d}' for m in re.finditer(r'(SM|UI)(\d+)(?:[—–-](?:SM|UI)?(\d+))?',str(v)) for n in range(int(m[2]),int(m[3] or m[2])+1)]
decl={r:ids(b.cell(r,3).value) for r in decision_rows}
cases={w[s].cell(r,1).value:(s,r,col) for s,n,col in [('冒烟用例',96,'V'),('UI检查',47,'AF')] for r in range(6,n+1)}
extra=[]
for cid,(s,cr,col) in cases.items():
    refs=set(map(int,re.findall(r"'本批提测'!\$H\$(\d+)",w[s][f'{col}{cr}'].value)))
    for r in refs:
        if cid not in decl.get(r,[]):extra.append({'case':cid,'decision':b.cell(r,1).value,'row':r})
before=[json.loads(l) for l in (root/'ui-color-live.ndjson').read_text(encoding='utf-8-sig').splitlines()]
now=[json.loads(l) for l in (root/'ui-color-recheck.ndjson').read_text(encoding='utf-8-sig').splitlines()]
assert before==now
result={'sourceUnchanged':True,'extraDependencies':extra,'SM82':[w['冒烟用例'].cell(87,c).value for c in [1,2,3,10,11]],'sourceMappings':[[w['UI检查'].cell(r,c).value for c in [1,18,28,36]] for r in range(6,48) if w['UI检查'].cell(r,28).value]}
(root/'color-independent-audit.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(result,ensure_ascii=False))
