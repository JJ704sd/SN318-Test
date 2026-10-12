import json, re
from pathlib import Path
import openpyxl

root=Path('./recipes/sn318-site-review')
out=Path('./artifacts/smoke-ui')
source=out/'SN318_提测前冒烟与UI规范_二次复核修正.xlsx'
target=out/'SN318_提测前冒烟与UI规范_网站对照修订.xlsx'
original=openpyxl.load_workbook(source,data_only=False)
book=openpyxl.load_workbook(target,data_only=False)
cached=openpyxl.load_workbook(target,data_only=True)
assert book.sheetnames==original.sheetnames+['网站详情对照']
checks=[]
for name,ranges in {
 '本批提测':['B7:B17','D40:G61','D115:G117','B33:B36'],
 '冒烟用例':['D6:H93','M6:R93'],
 'UI检查':['D6:H45','M6:V45']
}.items():
    for area in ranges:
        before=[[c.value for c in row] for row in original[name][area]]
        after=[[c.value for c in row] for row in book[name][area]]
        assert before==after,(name,area)
        checks.append(f'preserved {name} {area}')
case_map={}
for name,last,prefix,basis in [('冒烟用例',96,'SM','V'),('UI检查',46,'UI','AF')]:
    for r in range(6,last+1):
        row=book[name]
        cid=f'{prefix}{r-5:02d}'
        assert row.cell(r,1).value==cid
        assert row.cell(r,7).value=='未执行'
        assert isinstance(row[f'I{r}'].value,str) and row[f'I{r}'].value.startswith('=')
        case_map[cid]=(name,r,basis)
assert len(case_map)==132
assert sum(book['冒烟用例'].cell(r,16).value=='冒烟' for r in range(6,97))==28
assert sum(book['冒烟用例'].cell(r,16).value=='专项回归' for r in range(6,97))==63
for r in list(range(40,62))+list(range(115,118))+list(range(156,160)):
    clause=str(book['本批提测'].cell(r,3).value)
    for m in re.finditer(r'(SM|UI)(\d+)(?:[—–-](?:SM|UI)?(\d+))?',clause):
        for n in range(int(m[2]),int(m[3] or m[2])+1):
            cid=f'{m[1]}{n:02d}'
            sh,cr,col=case_map[cid]
            formula=book[sh][f'{col}{cr}'].value
            assert re.search(rf"'本批提测'!\$H\${r}(?!\d)",formula),(cid,r)
review_ids=lambda w:{w['UI检查'].cell(r,28).value for r in range(6,46) if w['UI检查'].cell(r,28).value}
assert review_ids(book)==review_ids(original) and len(review_ids(book))==20
assert [book['网站详情对照'].cell(r,1).value for r in range(6,16)]==[f'WEB{n:02d}' for n in range(1,11)]
table_refs={name:{t.name:t.ref for t in sheet.tables.values()} for name,sheet in [(s,book[s]) for s in book.sheetnames]}
assert 'A5:X96' in table_refs['冒烟用例'].values()
assert 'A5:AH46' in table_refs['UI检查'].values()
assert 'A5:H123' in table_refs['需求对照'].values()
assert 'A155:H159' in table_refs['本批提测'].values()
errors=[(sh.title,c.coordinate,c.value) for sh in cached for row in sh for c in row if c.data_type=='e']
assert not errors,errors
for cell in ['B22','F22','H22']:
    assert cached['本批提测'][cell].value==0,(cell,cached['本批提测'][cell].value)
assert cached['本批提测']['B6'].value=='未就绪'
for sh,areas in {'冒烟用例':['D6:H96','M6:R96'],'UI检查':['D6:H46','M6:V46'],'本批提测':['B7:B17','D40:G61','D115:G117','D156:G159']}.items():
    for area in areas:
        assert not any('复核模拟' in str(c.value) or c.value=='AUDIT-BUILD' for row in book[sh][area] for c in row),(sh,area)
result={'savedFile':str(target),'bytes':target.stat().st_size,'functional':91,'smoke':28,'regression':63,'ui':41,'allResults':'未执行','formulaErrors':0,'sourceReviewIdsPreserved':20,'decisionEdgesComplete':True,'tableRefs':table_refs,'preservedInputs':checks}
(root/'saved-file-verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(result,ensure_ascii=False))
