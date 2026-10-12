import json,re
from pathlib import Path
from collections import Counter
import openpyxl
root=Path('./recipes/sn318-site-review')
folder=Path('./artifacts/smoke-ui')
src=folder/'SN318_功能冒烟与界面交互验证_覆盖复核.xlsx'
dst=folder/'SN318_冒烟与UI用例_飞书五色及实际页面复核.xlsx'
before=openpyxl.load_workbook(src,data_only=False)
w=openpyxl.load_workbook(dst,data_only=False)
c=openpyxl.load_workbook(dst,data_only=True)
assert w.sheetnames==before.sheetnames
def values(sh):return {x.coordinate:x.value for row in sh for x in row if x.value is not None}
unchanged=['需求对照','UI修改清单','开发冒烟']
for name in unchanged:assert values(w[name])==values(before[name]),name
for name,areas in {'本批提测':['B7:B17','D40:G61','D115:G117','D156:G160','B33:B36'],'UI检查':['D6:I54','M6:O54','U6:V54'],'冒烟用例':['D6:I96','M6:O96'],'UI走查':['H7:Y55','A62:N70']}.items():
 for area in areas:
  assert [[x.value for x in row] for row in w[name][area]]==[[x.value for x in row] for row in before[name][area]],(name,area)
def clean(x):return re.sub(r'<at[^>]*>(.*?)</at>',r'\1','、'.join(x) if isinstance(x,list) else str(x or '')).strip()
live=[json.loads(l) for l in (root/'ui-color-page-recheck-live.ndjson').read_text(encoding='utf-8-sig').splitlines()]
valid={r['record_id']:r for r in live if any(clean(r.get(k)) for k in ['页面/模块','当前原型问题','Review 建议','后续处理建议'])}
mapped={w['UI检查'].cell(r,28).value:r for r in range(6,55) if w['UI检查'].cell(r,28).value}
assert set(mapped)==set(valid) and len(mapped)==21
counts=Counter()
palette={'无法覆盖':'FDE7E9','业务交互、信息待明确':'F3F4F6','组合复用':'E0F2FE','调整后复用':'FEF3C7','直接复用':'DCFCE7'}
ids={}
for rid,r in mapped.items():
 rec=valid[rid];level=re.sub(r'^[^\u4e00-\u9fff]+','',clean(rec['框架适配判断']))
 assert w['UI检查'].cell(r,18).value==level
 for col,key in [(16,'优先级'),(17,'确认状态'),(19,'当前处理角色'),(24,'Review 类型'),(25,'当前原型问题'),(26,'Review 建议'),(27,'后续处理建议')]:
  assert clean(w['UI检查'].cell(r,col).value)==clean(rec.get(key)),(rid,key)
 ident=w['UI检查'].cell(r,1).value;counts[level]+=1;ids[ident]=rid
assert counts=={'无法覆盖':3,'业务交互、信息待明确':1,'组合复用':2,'调整后复用':1,'直接复用':14}
assert {ident for ident,rid in ids.items() if w['UI检查'].cell(mapped[rid],18).value=='无法覆盖'}=={'UI01','UI02','UI42'}
assert {ident for ident,rid in ids.items() if w['UI检查'].cell(mapped[rid],18).value=='组合复用'}=={'UI03','UI35'}
assert len([r for r in range(6,55) if not w['UI检查'].cell(r,28).value])==28
audit={w['网站详情对照'].cell(r,1).value:r for r in range(24,45)}
assert set(audit)==set(ids) and len(audit)==21
for ident,r in audit.items():
 fr=mapped[ids[ident]];level=w['UI检查'].cell(fr,18).value
 assert w['网站详情对照'].cell(r,2).value==level
 assert ids[ident] in w['网站详情对照'].cell(r,3).value
 assert w['网站详情对照'][f'F{r}'].value==f"='UI检查'!AI{fr}"
 assert w['网站详情对照'][f'B{r}'].fill.fgColor.rgb.endswith(palette[level])
scope=[w['UI走查'][f'B{r}'].value for r in range(62,112) if w['UI走查'][f'B{r}'].value]
assert len(scope)==30 and set(ids)<=set(scope)
assert all(ident in scope for ident in ['UI33','UI34','UI48'])
assert len(set(w['UI走查'][f'A{r}'].value for r in range(62,92)))==30
for r in range(62,112):
 assert w['UI走查'][f'E{r}'].value=='待选择'
 assert w['UI走查'][f'H{r}'].value=='未执行'
 assert c['UI走查'][f'O{r}'].value==('未纳入' if r<92 else None)
 for col in ['O','P','Q','R']:assert w['UI走查'][f'{col}{r}'].value==before['UI走查'][f'{col}{r}'].value
for sh,last,prefix in [('冒烟用例',96,'SM'),('UI检查',54,'UI')]:
 assert [w[sh].cell(r,1).value for r in range(6,last+1)]==[f'{prefix}{r-5:02d}' for r in range(6,last+1)]
 assert all(w[sh].cell(r,7).value=='未执行' for r in range(6,last+1))
for a in ['B22','F22','H22']:assert c['本批提测'][a].value==0
assert c['本批提测']['C173'].value==22 and c['本批提测']['D173'].value==21
assert c['本批提测']['B6'].value=='未就绪'
assert c['开发冒烟']['E4'].value==c['UI走查']['E4'].value==c['UI修改清单']['E4'].value==0
assert '28' in w['本批提测']['B176'].value and '642' in w['本批提测']['B177'].value
assert '没有勾选列' in w['UI检查']['J18'].value
assert '默认/hover' in w['UI检查']['J19'].value
assert '取得专项布局' not in w['UI检查']['J40'].value and '必要时' in w['UI检查']['J40'].value
assert '加载完成' in w['UI检查']['J46'].value
assert '不作为本红色行内组件' in w['UI检查']['J47'].value
assert '未改字段' in w['UI检查']['J53'].value
assert [w['网站详情对照'].cell(r,1).value for r in range(16,19)]==['WEB11','WEB12','WEB13']
assert w['网站详情对照'].tables['SiteEvidence'].ref=='A5:H18'
assert w['网站详情对照'].tables['ReviewPageMapping'].ref=='A23:H44'
errors=[(sh.title,x.coordinate,x.value) for sh in c for row in sh for x in row if x.data_type=='e']
assert not errors,errors
report={'path':str(dst),'bytes':dst.stat().st_size,'reviewRevision':642,'prdRevision':8340,'uiSpecRevision':5256,'rawSourceRecords':22,'validMappings':21,'supplementaryCases':28,'counts':dict(counts),'actualPageMappingRows':21,'pageCandidates':30,'formalTests':140,'developerScenes':16,'allUnexecuted':True,'inputsAndResultFormulasPreserved':True,'formulaErrors':0,'unchangedSheets':unchanged}
(root/'color-page-saved-verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(report,ensure_ascii=False))
