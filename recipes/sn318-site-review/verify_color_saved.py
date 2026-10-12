import json, re
from pathlib import Path
from collections import Counter
import openpyxl

root=Path('./recipes/sn318-site-review')
folder=Path('./artifacts/smoke-ui')
original=openpyxl.load_workbook(folder/'SN318_提测前冒烟与UI规范_网站对照修订.xlsx',data_only=False)
path=folder/'SN318_提测前冒烟与UI规范_颜色层级对应_复查修正.xlsx'
book=openpyxl.load_workbook(path,data_only=False)
cached=openpyxl.load_workbook(path,data_only=True)
records=[json.loads(l) for l in (root/'ui-color-live.ndjson').read_text(encoding='utf-8-sig').splitlines()]
is_valid=lambda r:any(r.get(k) for k in ['页面/模块','当前原型问题','Review 建议','后续处理建议'])
valid={r['record_id']:r for r in records if is_valid(r)}
assert len(records)==22 and len(valid)==21
assert book.sheetnames==original.sheetnames
preserved=[]
for sh,ranges in {'本批提测':['B7:B17','D40:G61','D115:G117','D156:G159','B33:B36'],'冒烟用例':['D6:H96','M6:R96'],'UI检查':['D6:H46','M6:O46','U6:V46','J6:K46']}.items():
    for area in ranges:
        assert [[c.value for c in row] for row in book[sh][area]]==[[c.value for c in row] for row in original[sh][area]],(sh,area)
        preserved.append([sh,area])
for sh in ['网站详情对照']:
    assert [[c.value for c in row] for row in book[sh]]==[[c.value for c in row] for row in original[sh]],sh
case_map={}
for sh,last,prefix,col in [('冒烟用例',96,'SM','V'),('UI检查',47,'UI','AF')]:
    for r in range(6,last+1):
        cid=f'{prefix}{r-5:02d}'
        assert book[sh].cell(r,1).value==cid
        assert book[sh].cell(r,7).value=='未执行'
        case_map[cid]=(sh,r,col)
assert len(case_map)==133
mapped={book['UI检查'].cell(r,28).value:r for r in range(6,48) if book['UI检查'].cell(r,28).value}
assert set(mapped)==set(valid)
palette={'无法覆盖':'FDE7E9','组合复用':'E0F2FE','调整后复用':'FEF3C7','直接复用':'DCFCE7','业务交互、信息待明确':'F3F4F6'}
levels=Counter()
for rid,r in mapped.items():
    observed=re.sub(r'^[^\u4e00-\u9fff]+','',valid[rid]['框架适配判断'][0])
    assert book['UI检查'].cell(r,18).value==observed
    levels[observed]+=1
    for col in ['A','B','C','R']:
        assert book['UI检查'][f'{col}{r}'].fill.fgColor.rgb[-6:]==palette[observed],(rid,col)
    assert rid in book['UI检查'][f'AL{r}'].value
assert dict(levels)=={'无法覆盖':3,'组合复用':2,'调整后复用':1,'业务交互、信息待明确':1,'直接复用':14}
assert cached['本批提测']['C173'].value==22
assert cached['本批提测']['D173'].value==21
declared={cid:set() for cid in case_map}
for r in list(range(40,62))+list(range(115,118))+list(range(156,161)):
    clause=str(book['本批提测'].cell(r,3).value)
    for m in re.finditer(r'(SM|UI)(\d+)(?:[—–-](?:SM|UI)?(\d+))?',clause):
        for n in range(int(m[2]),int(m[3] or m[2])+1):
            cid=f'{m[1]}{n:02d}'
            sh,cr,col=case_map[cid]
            declared[cid].add(r)
            assert re.search(rf"'本批提测'!\$H\${r}(?!\d)",book[sh][f'{col}{cr}'].value),(cid,r)
for cid,(sh,r,col) in case_map.items():
    actual=set(map(int,re.findall(r"'本批提测'!\$H\$(\d+)",book[sh][f'{col}{r}'].value)))
    assert actual==declared[cid],(cid,actual-declared[cid],declared[cid]-actual)
table_refs={sh:{t.name:t.ref for t in book[sh].tables.values()} for sh in book.sheetnames}
assert 'A5:AM47' in table_refs['UI检查'].values()
assert 'A5:Z96' in table_refs['冒烟用例'].values()
assert 'A155:H160' in table_refs['本批提测'].values()
assert 'A5:H124' in table_refs['需求对照'].values()
errors=[(sh.title,c.coordinate,c.value) for sh in cached for row in sh for c in row if c.data_type=='e']
assert not errors,errors
for cell in ['B22','F22','H22']:
    assert cached['本批提测'][cell].value==0
assert cached['本批提测']['B6'].value=='未就绪'
review=json.loads((root/'color-repaired-verification.json').read_text(encoding='utf-8'))
reverse={}
for row in review['perCase']:
    expected='、'.join(row['related']) or '无直接配套；本UI case独立执行'
    assert book['UI检查'][f"AJ{row['row']}"].value==expected
    assert book['UI检查'][f"AM{row['row']}"].value==row['scope']
    for cid in row['related']:reverse.setdefault(cid,[]).append(row['id'])
for cid,(sh,r,col) in case_map.items():
    if sh=='冒烟用例':assert book[sh][f'Y{r}'].value==('、'.join(reverse.get(cid,[])) or None),cid
prior=openpyxl.load_workbook(folder/'SN318_提测前冒烟与UI规范_颜色层级对应.xlsx',data_only=False)
for sh,last in [('冒烟用例',96),('UI检查',47)]:
    for r in range(6,last+1):
        assert book[sh][f'I{r}'].value==prior[sh][f'I{r}'].value,(sh,r)
result={'path':str(path),'bytes':path.stat().st_size,'sourceRecords':22,'mappedSourceRecords':21,'sourceColorCaseCounts':dict(levels),'totalCases':133,'allExecutionResults':'未执行','formulaErrors':0,'decisionEdgesExact':True,'functionalLinksBidirectional':True,'nativeColorCellsVerified':True,'executionInputsPreserved':preserved,'tableRefs':table_refs}
(root/'color-rechecked-saved-verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(result,ensure_ascii=False))
