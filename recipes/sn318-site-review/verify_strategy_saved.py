import json,re
from pathlib import Path
from collections import Counter
import openpyxl
root=Path('./recipes/sn318-site-review')
folder=Path('./artifacts/smoke-ui')
src=folder/'SN318_提测前冒烟与UI规范_颜色层级对应_复查修正.xlsx'
dst=folder/'SN318_功能冒烟与界面交互验证_策略调整.xlsx'
before=openpyxl.load_workbook(src,data_only=False)
w=openpyxl.load_workbook(dst,data_only=False)
c=openpyxl.load_workbook(dst,data_only=True)
assert w.sheetnames==before.sheetnames+['开发冒烟','UI走查','UI修改清单']
preserved=[]
for name,areas in {'本批提测':['B7:B17','D40:G61','D115:G117','D156:G160','B33:B36'],'冒烟用例':['D6:H96','M6:O96','Q6:R96'],'UI检查':['D6:H47','M6:O47','U6:V47']}.items():
    for area in areas:
        assert [[x.value for x in row] for row in w[name][area]]==[[x.value for x in row] for row in before[name][area]],(name,area)
        preserved.append([name,area])
case_map={}
for name,last,prefix,basis in [('冒烟用例',96,'SM','V'),('UI检查',54,'UI','AF')]:
    for r in range(6,last+1):
        ident=f'{prefix}{r-5:02d}'
        assert w[name].cell(r,1).value==ident
        assert w[name].cell(r,7).value=='未执行'
        case_map[ident]=(name,r,basis)
        if r <= (96 if name=='冒烟用例' else 47):
            assert w[name][f'I{r}'].value==before[name][f'I{r}'].value
assert len(case_map)==140
assert Counter(w['冒烟用例'].cell(r,16).value for r in range(6,97))=={'主流程回归':28,'专项回归':63}
declared={ident:set() for ident in case_map}
for dr in list(range(40,62))+list(range(115,118))+list(range(156,161)):
    for m in re.finditer(r'(SM|UI)(\d+)(?:[—–-](?:SM|UI)?(\d+))?',str(w['本批提测'].cell(dr,3).value)):
        for n in range(int(m[2]),int(m[3] or m[2])+1):declared[f'{m[1]}{n:02d}'].add(dr)
for ident,(sh,r,col) in case_map.items():
    refs=set(map(int,re.findall(r"'本批提测'!\$H\$(\d+)",w[sh][f'{col}{r}'].value)))
    assert refs==declared[ident],(ident,refs-declared[ident],declared[ident]-refs)
live=[json.loads(l) for l in (root/'ui-strategy-live.ndjson').read_text(encoding='utf-8-sig').splitlines()]
valid={r['record_id']:r for r in live if any(r.get(k) for k in ['页面/模块','当前原型问题','Review 建议','后续处理建议'])}
mapped={w['UI检查'].cell(r,28).value:r for r in range(6,55) if w['UI检查'].cell(r,28).value}
assert set(mapped)==set(valid) and len(mapped)==21
for rid,r in mapped.items():
    assert w['UI检查'].cell(r,17).value=='、'.join(valid[rid]['确认状态'])
    assert w['UI检查'].cell(r,18).value==re.sub(r'^[^\u4e00-\u9fff]+','',valid[rid]['框架适配判断'][0])
assert w['UI检查']['Q6'].value==w['UI检查']['Q7'].value=='设计完毕'
for n in range(1,50):
    r=n+6;ident=f'UI{n:02d}'
    assert w['UI走查'].cell(r,1).value==ident
    assert w['UI走查'].cell(r,9).value=='未走查'
    assert w['UI走查'][f'F{r}'].value==f"='UI检查'!J{n+5}"
    assert w['UI走查'][f'G{r}'].value==f"='UI检查'!K{n+5}"
    assert w['UI走查'][f'T{r}'].value==f"='UI检查'!I{n+5}"
for n in range(1,17):
    r=n+6
    assert w['开发冒烟'].cell(r,1).value==f'DS{n:02d}'
    assert w['开发冒烟'].cell(r,9).value=='未执行'
    assert '$B$11' not in w['开发冒烟'][f'N{r}'].value
assert '131' in w['开发冒烟']['E16'].value and '20KG' in w['开发冒烟']['E16'].value
for r in range(7,28):
    assert w['UI修改清单'].cell(r,4).value=='Review待核对'
    assert w['UI修改清单'].cell(r,14).value=='待走查核对'
    assert c['UI修改清单'][f'U{r}'].value=='待走查候选'
assert c['开发冒烟']['E4'].value==0
assert c['UI走查']['E4'].value==0
assert c['UI修改清单']['B4'].value==21
assert c['UI修改清单']['E4'].value==0
assert c['UI修改清单']['H4'].value==0
for cell in ['B22','F22','H22']:assert c['本批提测'][cell].value==0
assert c['本批提测']['C173'].value==22 and c['本批提测']['D173'].value==21
assert c['本批提测']['B6'].value=='未就绪'
errors=[(sh.title,x.coordinate,x.value) for sh in c for row in sh for x in row if x.data_type=='e']
assert not errors,errors
for sh,areas in {'本批提测':['B7:B17','D40:G61','D115:G117','D156:G160'],'开发冒烟':['G7:M22'],'UI走查':['H7:R55'],'UI修改清单':['D7:S37']}.items():
    assert not any('复核模拟' in str(x.value) or x.value=='AUDIT-BUILD' for area in areas for row in w[sh][area] for x in row),(sh,area)
refs={sh:{t.name:t.ref for t in w[sh].tables.values()} for sh in w.sheetnames}
assert 'A5:AM54' in refs['UI检查'].values()
assert 'A5:AB96' in refs['冒烟用例'].values()
assert 'A6:P22' in refs['开发冒烟'].values()
assert 'A6:T55' in refs['UI走查'].values()
assert 'A6:U37' in refs['UI修改清单'].values()
assert 'A5:H131' in refs['需求对照'].values()
report={'path':str(dst),'bytes':dst.stat().st_size,'sourceReviewRevision':642,'publicBaselineRevision':95,'formalCases':140,'functional':91,'ui':49,'developerScenes':16,'reviewCandidates':21,'actualDefects':0,'allTestsUnexecuted':True,'walkthroughUnstarted':True,'exactDecisionEdges':True,'sourceIdsMatched':21,'inputsPreserved':preserved,'formulaErrors':0,'tableRefs':refs}
(root/'strategy-saved-verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(report,ensure_ascii=False))
